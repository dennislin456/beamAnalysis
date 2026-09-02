import os
import json
import pandas as pd
import numpy as np
import pyqtgraph as pg
from scipy.ndimage import uniform_filter, map_coordinates

from PyQt5.QtWidgets import (QWidget, QHBoxLayout, QVBoxLayout, QPushButton, 
                             QLabel, QFileDialog, QMessageBox, QFrame, 
                             QSplitter, QSizePolicy, QRadioButton, QButtonGroup, 
                             QScrollArea, QCheckBox, QComboBox, QApplication)
from PyQt5.QtCore import Qt, QTimer

import openpyxl
from openpyxl.chart import LineChart, Reference

from batch_data_loader import load_numeric_matrix

# 💡 這裡將匯入共用的元件 (之後會統一放在 shared_components.py 中)
from shared_components import (NoWheelSpinBox, NoWheelDoubleSpinBox,
                               HeatmapViewerWindow, CrossProfileViewerWindow,
                               compute_auto_spot_center, build_robust_threshold_mask,
                               split_y_index, apply_readable_plot_theme,
                               LevelAlignedHistogramLUTItem,
                               export_plot_image, EXPORT_IMAGE_EXT)

class DataRayTab(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        
        # --- DataRay 專屬變數初始化 ---
        self.save_dir_path = ""
        self.file1_path = ""
        self.file2_path = ""
        self.param_file_path = ""
        self.matrix1 = None
        self.matrix2 = None
        self.result_matrix = None
        self.smoothed_matrix = None
        
        self.click_points = []
        self.measure_items = []
        self.contour_click_points = []
        self.contour_measure_items = []
        self.contour_curves = []

        self.dataray_center = None 
        self.dataray_circle_item = None
        self.dataray_center_spot = None
        self.dataray_fixed_cross_items = []
        self.is_updating_dataray_ui = False

        self.heatmap_cross_point = None
        self.heatmap_cross_items = []
        self.m1_center_point = None
        self.m1_marker_items = []

        self.viewer_m1_win = None
        self.viewer_m2_win = None
        self.cross_profile_win = None
        
        # 定時器初始化
        self.ma_timer = QTimer()
        self.ma_timer.setSingleShot(True)
        self.ma_timer.timeout.connect(self.apply_ma_kernel_change)
        
        self.cbar_timer = QTimer()
        self.cbar_timer.setSingleShot(True)
        self.cbar_timer.timeout.connect(self.update_isocurves_and_waveform)

        # 建立 UI
        self.setup_ui()

    def setup_ui(self):
        btn_style_default = """
            QPushButton { font-size: 13px; font-weight: bold; background-color: #f0f0f0; border: 1px solid #cccccc; border-radius: 5px; padding: 6px 12px; }
            QPushButton:hover { background-color: #e0e0e0; border-color: #b0b0b0; }
            QPushButton:pressed { background-color: #d0d0d0; }
        """
        btn_style_folder = """
            QPushButton { font-size: 13px; font-weight: bold; color: white; background-color: #EF6C00; border: none; border-radius: 5px; padding: 6px 12px; }
            QPushButton:hover { background-color: #F57C00; }
            QPushButton:pressed { background-color: #E65100; }
        """
        btn_style_view = """
            QPushButton { font-size: 13px; font-weight: bold; color: white; background-color: #7B1FA2; border: none; border-radius: 5px; padding: 6px 14px; }
            QPushButton:hover { background-color: #8E24AA; }
            QPushButton:pressed { background-color: #4A148C; }
            QPushButton:disabled { background-color: #E0E0E0; color: #A0A0A0; }
        """
        btn_style_cross = """
            QPushButton { font-size: 13px; font-weight: bold; color: white; background-color: #00897B; border: none; border-radius: 5px; padding: 6px 14px; }
            QPushButton:hover { background-color: #009688; }
            QPushButton:pressed { background-color: #00695C; }
            QPushButton:disabled { background-color: #E0E0E0; color: #A0A0A0; }
        """
        btn_style_primary = """
            QPushButton { font-size: 13px; font-weight: bold; color: white; background-color: #2E7D32; border: none; border-radius: 5px; padding: 8px 12px; }
            QPushButton:hover { background-color: #388E3C; }
            QPushButton:pressed { background-color: #1B5E20; }
        """
        btn_style_export = """
            QPushButton { font-size: 13px; font-weight: bold; color: white; background-color: #0288D1; border: none; border-radius: 5px; padding: 8px 16px; }
            QPushButton:hover { background-color: #039BE5; }
            QPushButton:pressed { background-color: #01579B; }
            QPushButton:disabled { background-color: #B0BEC5; }
        """
        btn_style_danger = """
            QPushButton { font-size: 13px; font-weight: bold; color: white; background-color: #D32F2F; border: none; border-radius: 5px; padding: 6px 12px; }
            QPushButton:hover { background-color: #E53935; }
            QPushButton:pressed { background-color: #C62828; }
        """

        main_layout = QHBoxLayout(self)
        main_layout.setContentsMargins(6, 6, 6, 6)
        
        splitter = QSplitter(Qt.Horizontal)
        splitter.setStyleSheet("QSplitter::handle { background-color: #dcdcdc; width: 4px; }")
        main_layout.addWidget(splitter)
        
        # --- 左側主容器 ---
        left_container = QWidget()
        left_container.setMinimumWidth(320)
        left_container.setMaximumWidth(360)
        left_outer_layout = QVBoxLayout(left_container)
        left_outer_layout.setContentsMargins(0, 0, 0, 0)
        left_outer_layout.setSpacing(0)

        # 上方固定面板
        top_fixed_widget = QWidget()
        top_fixed_layout = QVBoxLayout(top_fixed_widget)
        top_fixed_layout.setContentsMargins(12, 12, 12, 12)
        top_fixed_layout.setSpacing(6)
        
        lbl_mode_title = QLabel("選擇工作模式：")
        lbl_mode_title.setStyleSheet("font-weight: bold; color: #333333; margin-top: 2px;")
        top_fixed_layout.addWidget(lbl_mode_title)
        
        self.combo_mode = QComboBox()
        self.combo_mode.addItem("雙檔峰值校正 (R = M1 - k * M2)", "calc")
        self.combo_mode.addItem("雙檔純相減畫圖 (M1 - M2)", "sub")
        self.combo_mode.addItem("雙檔純相除畫圖 (M1 / M2)", "div")
        self.combo_mode.addItem("單獨匯入畫圖", "single")
        self.combo_mode.addItem("匯入數據檔 + 載入參數檔重繪", "param")
        self.combo_mode.currentIndexChanged.connect(self.on_mode_changed)
        top_fixed_layout.addWidget(self.combo_mode)

        self.chk_normalize_peaks = QCheckBox("Normalize")
        self.chk_normalize_peaks.setChecked(False)
        self.chk_normalize_peaks.setStyleSheet("font-weight: bold; color: #D81B60;")
        top_fixed_layout.addWidget(self.chk_normalize_peaks)
        
        top_fixed_layout.addWidget(self._create_hline())
        
        self.btn_file1 = QPushButton("I. 匯入第一個檔案")
        self.btn_file1.setStyleSheet(btn_style_default)
        self.btn_file1.clicked.connect(self.load_file1)
        top_fixed_layout.addWidget(self.btn_file1)
        
        self.lbl_file1 = QLabel("未選擇檔案")
        self.lbl_file1.setStyleSheet("color: #757575; font-size: 11px;")
        self.lbl_file1.setWordWrap(True)
        top_fixed_layout.addWidget(self.lbl_file1)
        
        self.btn_file2 = QPushButton("II. 匯入第二個檔案")
        self.btn_file2.setStyleSheet(btn_style_default)
        self.btn_file2.clicked.connect(self.load_file2)
        top_fixed_layout.addWidget(self.btn_file2)
        
        self.lbl_file2 = QLabel("未選擇檔案")
        self.lbl_file2.setStyleSheet("color: #757575; font-size: 11px;")
        self.lbl_file2.setWordWrap(True)
        top_fixed_layout.addWidget(self.lbl_file2)
        
        top_fixed_layout.addWidget(self._create_hline())
        
        self.btn_calc = QPushButton("開始計算與畫圖")
        self.btn_calc.setStyleSheet(btn_style_primary)
        self.btn_calc.setMinimumHeight(38)
        self.btn_calc.clicked.connect(self.process_data)
        top_fixed_layout.addWidget(self.btn_calc)
        
        top_fixed_layout.addWidget(self._create_hline())

        lbl_status_title = QLabel("計算數據與狀態")
        lbl_status_title.setStyleSheet("font-weight: bold; font-size: 12px; color: #37474F;")
        top_fixed_layout.addWidget(lbl_status_title)
        
        self.lbl_status = QLabel("狀態: 等待匯入檔案")
        self.lbl_status.setStyleSheet("color: #1565C0; font-weight: bold; font-size: 12px;")
        top_fixed_layout.addWidget(self.lbl_status)

        self.lbl_max1 = QLabel("M1 最大值(位置): --")
        self.lbl_max1.setStyleSheet("font-size: 12px;")
        self.lbl_max2 = QLabel("M2 同位置數值: --")
        self.lbl_max2.setStyleSheet("font-size: 12px;")
        self.lbl_ratio = QLabel("計算得出比例: --")
        self.lbl_ratio.setStyleSheet("font-size: 12px;")
        self.lbl_size = QLabel("矩陣大小: --")
        self.lbl_size.setStyleSheet("font-size: 12px;")
        
        top_fixed_layout.addWidget(self.lbl_max1)
        top_fixed_layout.addWidget(self.lbl_max2)
        top_fixed_layout.addWidget(self.lbl_ratio)
        top_fixed_layout.addWidget(self.lbl_size)
        
        self.lbl_distance = QLabel("位置差距: ΔX: --, ΔY: -- \n總距離: -- px")
        self.lbl_distance.setStyleSheet("font-weight: bold; color: #37474F; font-size: 12px;")
        top_fixed_layout.addWidget(self.lbl_distance)
        
        self.lbl_real_distance = QLabel("實際差距 (* 5.5): ΔX: -- μm, ΔY: -- μm \n總距離: -- μm")
        self.lbl_real_distance.setStyleSheet("font-weight: bold; color: #2E7D32; font-size: 12px;")
        top_fixed_layout.addWidget(self.lbl_real_distance)

        left_outer_layout.addWidget(top_fixed_widget)

        # 下方捲動面板
        left_scroll_panel = QScrollArea()
        left_scroll_panel.setWidgetResizable(True)
        
        left_content = QWidget()
        left_layout = QVBoxLayout(left_content)
        left_layout.setContentsMargins(12, 12, 12, 12)
        left_layout.setSpacing(10)

        left_layout.addWidget(self._create_hline())

        lbl_center_mode = QLabel("M1 / 雙點定位模式：")
        lbl_center_mode.setStyleSheet("font-weight: bold;")
        left_layout.addWidget(lbl_center_mode)

        lbl_m1_pt = QLabel("第一點（M1 Heatmap）：")
        lbl_m1_pt.setStyleSheet("font-weight: bold; color: #2E7D32;")
        left_layout.addWidget(lbl_m1_pt)

        self.m1_point_group = QButtonGroup(self)
        self.radio_m1_centroid = QRadioButton("自動抓取 (質心中心)")
        self.radio_m1_centroid.setChecked(True)
        self.radio_m1_centroid.toggled.connect(self.on_m1_point_mode_changed)
        self.m1_point_group.addButton(self.radio_m1_centroid)
        left_layout.addWidget(self.radio_m1_centroid)

        self.radio_m1_thresh_geom = QRadioButton("自動抓取 (門檻區域幾何中心)")
        self.radio_m1_thresh_geom.toggled.connect(self.on_m1_point_mode_changed)
        self.m1_point_group.addButton(self.radio_m1_thresh_geom)
        left_layout.addWidget(self.radio_m1_thresh_geom)

        self.radio_m1_peak_geom = QRadioButton("自動抓取 (最高值區域幾何中心)")
        self.radio_m1_peak_geom.toggled.connect(self.on_m1_point_mode_changed)
        self.m1_point_group.addButton(self.radio_m1_peak_geom)
        left_layout.addWidget(self.radio_m1_peak_geom)

        self.radio_m1_manual = QRadioButton("手動抓取 (點擊 M1 影像)")
        self.radio_m1_manual.toggled.connect(self.on_m1_point_mode_changed)
        self.m1_point_group.addButton(self.radio_m1_manual)
        left_layout.addWidget(self.radio_m1_manual)

        self.chk_m1_use_threshold = QCheckBox("使用門檻（第一點）")
        self.chk_m1_use_threshold.setChecked(True)
        self.chk_m1_use_threshold.setStyleSheet("color: #2E7D32;")
        self.chk_m1_use_threshold.toggled.connect(self.on_m1_threshold_toggled)
        left_layout.addWidget(self.chk_m1_use_threshold)

        layout_m1_thresh = QHBoxLayout()
        self.lbl_m1_thresh_spin = QLabel("第一點門檻比例 (%):")
        self.spin_m1_thresh_percent = NoWheelDoubleSpinBox()
        self.spin_m1_thresh_percent.setRange(0.1, 100.0)
        self.spin_m1_thresh_percent.setValue(50.0)
        self.spin_m1_thresh_percent.setSingleStep(1.0)
        self.spin_m1_thresh_percent.setDecimals(1)
        self.spin_m1_thresh_percent.valueChanged.connect(self.on_m1_thresh_percent_changed)
        layout_m1_thresh.addWidget(self.lbl_m1_thresh_spin)
        layout_m1_thresh.addWidget(self.spin_m1_thresh_percent)
        left_layout.addLayout(layout_m1_thresh)

        self.chk_m1_show_thresh = QCheckBox("顯示門檻區域於 M1 圖（洋紅半透明）")
        self.chk_m1_show_thresh.setChecked(True)
        self.chk_m1_show_thresh.setStyleSheet("color: #2E7D32;")
        self.chk_m1_show_thresh.toggled.connect(self.update_m1_thresh_overlay)
        left_layout.addWidget(self.chk_m1_show_thresh)

        lbl_p2_pt = QLabel("第二點（Process Result Heatmap）：")
        lbl_p2_pt.setStyleSheet("font-weight: bold; color: #1565C0;")
        left_layout.addWidget(lbl_p2_pt)

        self.p2_point_group = QButtonGroup(self)
        self.radio_p2_auto_min = QRadioButton("自動抓取 (第一點 Y 以下最小值)")
        self.radio_p2_auto_min.setChecked(True)
        self.radio_p2_auto_min.toggled.connect(self.on_p2_point_mode_changed)
        self.p2_point_group.addButton(self.radio_p2_auto_min)
        left_layout.addWidget(self.radio_p2_auto_min)

        self.radio_p2_m2_thresh_geom = QRadioButton("自動抓取 (M2 第一點 Y 以下門檻幾何中心)")
        self.radio_p2_m2_thresh_geom.toggled.connect(self.on_p2_point_mode_changed)
        self.p2_point_group.addButton(self.radio_p2_m2_thresh_geom)
        left_layout.addWidget(self.radio_p2_m2_thresh_geom)

        self.radio_p2_m2_centroid = QRadioButton("自動抓取 (M2 第一點 Y 以下質心中心)")
        self.radio_p2_m2_centroid.toggled.connect(self.on_p2_point_mode_changed)
        self.p2_point_group.addButton(self.radio_p2_m2_centroid)
        left_layout.addWidget(self.radio_p2_m2_centroid)

        self.radio_p2_manual = QRadioButton("手動抓取 (點擊 Process 影像)")
        self.radio_p2_manual.toggled.connect(self.on_p2_point_mode_changed)
        self.p2_point_group.addButton(self.radio_p2_manual)
        left_layout.addWidget(self.radio_p2_manual)

        self.chk_p2_use_threshold = QCheckBox("使用門檻（第二點／M2）")
        self.chk_p2_use_threshold.setChecked(True)
        self.chk_p2_use_threshold.setStyleSheet("color: #1565C0;")
        self.chk_p2_use_threshold.toggled.connect(self.on_p2_threshold_toggled)
        left_layout.addWidget(self.chk_p2_use_threshold)

        layout_p2_thresh = QHBoxLayout()
        self.lbl_p2_thresh_spin = QLabel("第二點門檻比例 (%):")
        self.spin_p2_thresh_percent = NoWheelDoubleSpinBox()
        self.spin_p2_thresh_percent.setRange(0.1, 100.0)
        self.spin_p2_thresh_percent.setValue(50.0)
        self.spin_p2_thresh_percent.setSingleStep(1.0)
        self.spin_p2_thresh_percent.setDecimals(1)
        self.spin_p2_thresh_percent.valueChanged.connect(self.on_p2_thresh_percent_changed)
        layout_p2_thresh.addWidget(self.lbl_p2_thresh_spin)
        layout_p2_thresh.addWidget(self.spin_p2_thresh_percent)
        left_layout.addLayout(layout_p2_thresh)

        lbl_p2_algo_hint = QLabel("質心／門檻幾何：背景扣除＋最大連通區＋亞像素（預設門檻 50%）")
        lbl_p2_algo_hint.setStyleSheet("color: #546E7A; font-size: 11px;")
        lbl_p2_algo_hint.setWordWrap(True)
        left_layout.addWidget(lbl_p2_algo_hint)

        self.chk_p2_show_thresh = QCheckBox("顯示門檻區域於 M2 圖（洋紅半透明，限 Y 以下）")
        self.chk_p2_show_thresh.setChecked(True)
        self.chk_p2_show_thresh.setStyleSheet("color: #1565C0;")
        self.chk_p2_show_thresh.toggled.connect(self.update_m2_thresh_overlay)
        left_layout.addWidget(self.chk_p2_show_thresh)

        left_layout.addWidget(self._create_hline())

        self.chk_enable_heatmap_cross = QCheckBox("啟用點擊定格十字標記")
        self.chk_enable_heatmap_cross.setChecked(True)
        self.chk_enable_heatmap_cross.setStyleSheet("font-weight: bold; color: #00897B;")
        self.chk_enable_heatmap_cross.toggled.connect(self.on_heatmap_cross_toggled)
        left_layout.addWidget(self.chk_enable_heatmap_cross)

        self.chk_enable_measure_cross = QCheckBox("啟用兩點距離量測十字標記")
        self.chk_enable_measure_cross.setChecked(False)
        self.chk_enable_measure_cross.setStyleSheet("font-weight: bold; color: #E65100;")
        self.chk_enable_measure_cross.toggled.connect(self.on_measure_cross_toggled)
        left_layout.addWidget(self.chk_enable_measure_cross)

        left_layout.addWidget(self._create_hline())

        lbl_dr_shape_title = QLabel("抓取光斑形狀選項：")
        lbl_dr_shape_title.setStyleSheet("font-weight: bold; color: #1565C0;")
        left_layout.addWidget(lbl_dr_shape_title)

        self.chk_dr_enable_spot = QCheckBox("啟用光斑抓取（隨第一點移動）")
        self.chk_dr_enable_spot.setChecked(False)
        self.chk_dr_enable_spot.setStyleSheet("font-weight: bold; color: #2E7D32;")
        self.chk_dr_enable_spot.toggled.connect(self.on_dr_enable_spot_toggled)
        left_layout.addWidget(self.chk_dr_enable_spot)

        self.dr_shape_group = QButtonGroup(self)

        layout_dr_shape_radios = QHBoxLayout()
        self.radio_dr_shape_circle = QRadioButton("正圓 (Circle)")
        self.radio_dr_shape_circle.setChecked(True)
        self.radio_dr_shape_circle.toggled.connect(self.on_dataray_shape_type_changed)

        self.radio_dr_shape_ellipse = QRadioButton("橢圓 (Ellipse)")
        self.radio_dr_shape_ellipse.toggled.connect(self.on_dataray_shape_type_changed)

        self.dr_shape_group.addButton(self.radio_dr_shape_circle)
        self.dr_shape_group.addButton(self.radio_dr_shape_ellipse)

        layout_dr_shape_radios.addWidget(self.radio_dr_shape_circle)
        layout_dr_shape_radios.addWidget(self.radio_dr_shape_ellipse)
        left_layout.addLayout(layout_dr_shape_radios)

        self.container_dr_circle_spin = QWidget()
        layout_dr_circle_spin = QHBoxLayout(self.container_dr_circle_spin)
        layout_dr_circle_spin.setContentsMargins(0, 0, 0, 0)
        lbl_dr_dia = QLabel("光斑直徑 Circle Size (px):")
        self.spin_dr_circle_diameter = NoWheelSpinBox()
        self.spin_dr_circle_diameter.setRange(2, 5000)
        self.spin_dr_circle_diameter.setValue(40)
        self.spin_dr_circle_diameter.setSingleStep(5)
        self.spin_dr_circle_diameter.valueChanged.connect(self.update_dataray_circle)
        layout_dr_circle_spin.addWidget(lbl_dr_dia)
        layout_dr_circle_spin.addWidget(self.spin_dr_circle_diameter)
        left_layout.addWidget(self.container_dr_circle_spin)

        self.container_dr_ellipse_spin = QWidget()
        layout_dr_ellipse_spin = QVBoxLayout(self.container_dr_ellipse_spin)
        layout_dr_ellipse_spin.setContentsMargins(0, 0, 0, 0)

        layout_dr_ellipse_x = QHBoxLayout()
        lbl_dr_ellipse_x = QLabel("X 軸直徑 Wx (px):")
        self.spin_dr_ellipse_wx = NoWheelSpinBox()
        self.spin_dr_ellipse_wx.setRange(2, 5000)
        self.spin_dr_ellipse_wx.setValue(40)
        self.spin_dr_ellipse_wx.setSingleStep(5)
        self.spin_dr_ellipse_wx.valueChanged.connect(self.update_dataray_circle)
        layout_dr_ellipse_x.addWidget(lbl_dr_ellipse_x)
        layout_dr_ellipse_x.addWidget(self.spin_dr_ellipse_wx)

        layout_dr_ellipse_y = QHBoxLayout()
        lbl_dr_ellipse_y = QLabel("Y 軸直徑 Wy (px):")
        self.spin_dr_ellipse_wy = NoWheelSpinBox()
        self.spin_dr_ellipse_wy.setRange(2, 5000)
        self.spin_dr_ellipse_wy.setValue(40)
        self.spin_dr_ellipse_wy.setSingleStep(5)
        self.spin_dr_ellipse_wy.valueChanged.connect(self.update_dataray_circle)
        layout_dr_ellipse_y.addWidget(lbl_dr_ellipse_y)
        layout_dr_ellipse_y.addWidget(self.spin_dr_ellipse_wy)

        layout_dr_ellipse_spin.addLayout(layout_dr_ellipse_x)
        layout_dr_ellipse_spin.addLayout(layout_dr_ellipse_y)
        self.container_dr_ellipse_spin.setVisible(False)
        left_layout.addWidget(self.container_dr_ellipse_spin)

        self.chk_dr_use_threshold = QCheckBox("使用門檻計算光斑寬度")
        self.chk_dr_use_threshold.setChecked(True)
        self.chk_dr_use_threshold.setStyleSheet("font-weight: bold; color: #C62828;")
        self.chk_dr_use_threshold.toggled.connect(self.on_dataray_threshold_toggled)
        left_layout.addWidget(self.chk_dr_use_threshold)

        layout_dr_thresh = QHBoxLayout()
        self.lbl_dr_thresh_spin = QLabel("光斑門檻比例 (%):")
        self.spin_dr_thresh_percent = NoWheelDoubleSpinBox()
        self.spin_dr_thresh_percent.setRange(0.1, 100.0)
        self.spin_dr_thresh_percent.setValue(13.5)
        self.spin_dr_thresh_percent.setSingleStep(1.0)
        self.spin_dr_thresh_percent.setDecimals(1)
        self.spin_dr_thresh_percent.valueChanged.connect(self.on_dr_thresh_percent_changed)
        layout_dr_thresh.addWidget(self.lbl_dr_thresh_spin)
        layout_dr_thresh.addWidget(self.spin_dr_thresh_percent)
        left_layout.addLayout(layout_dr_thresh)

        self.chk_dr_show_thresh = QCheckBox("顯示光斑門檻區域於 Process 圖（洋紅半透明）")
        self.chk_dr_show_thresh.setChecked(False)
        self.chk_dr_show_thresh.setStyleSheet("color: #C62828;")
        self.chk_dr_show_thresh.toggled.connect(self.update_process_thresh_overlay)
        left_layout.addWidget(self.chk_dr_show_thresh)

        left_layout.addWidget(self._create_hline())

        lbl_measure_title = QLabel("距離量測與分色十字工具")
        lbl_measure_title.setStyleSheet("font-weight: bold; font-size: 13px; color: #C62828;")
        left_layout.addWidget(lbl_measure_title)
        
        self.lbl_cursor = QLabel("目前滑鼠： X: -- , Y: --")
        self.lbl_cursor.setStyleSheet("color: #616161;")
        left_layout.addWidget(self.lbl_cursor)
        
        self.lbl_measure_points = QLabel("點擊點 1 (黃): --\n點擊點 2 (青藍): --")
        left_layout.addWidget(self.lbl_measure_points)
        
        layout_cross_ctrl = QHBoxLayout()
        lbl_cross_size = QLabel("標記十字大小 (px):")
        self.spin_cross_size = NoWheelSpinBox()
        self.spin_cross_size.setRange(10, 200)
        self.spin_cross_size.setValue(40)
        self.spin_cross_size.valueChanged.connect(self.redraw_measure_crosses)
        layout_cross_ctrl.addWidget(lbl_cross_size)
        layout_cross_ctrl.addWidget(self.spin_cross_size)
        left_layout.addLayout(layout_cross_ctrl)
        
        self.btn_clear_points = QPushButton("清除量測點與十字")
        self.btn_clear_points.setStyleSheet(btn_style_default)
        self.btn_clear_points.clicked.connect(self.clear_measure_points)
        left_layout.addWidget(self.btn_clear_points)
        
        left_layout.addWidget(self._create_hline())
        
        lbl_ma_title = QLabel("Contour 圖像設定")
        lbl_ma_title.setStyleSheet("font-weight: bold; font-size: 13px; color: #6A1B9A;")
        left_layout.addWidget(lbl_ma_title)
        
        layout_ma_ctrl = QHBoxLayout()
        lbl_ma_size = QLabel("Moving Average Kernel (px) :")
        self.spin_ma_size = NoWheelSpinBox()
        self.spin_ma_size.setRange(1, 99)
        self.spin_ma_size.setSingleStep(2)
        self.spin_ma_size.setValue(31)
        self.spin_ma_size.valueChanged.connect(self.on_ma_kernel_spin_changed)
        
        layout_ma_ctrl.addWidget(lbl_ma_size)
        layout_ma_ctrl.addWidget(self.spin_ma_size)
        left_layout.addLayout(layout_ma_ctrl)
        
        self.btn_clear_contour_lines = QPushButton("清除 Contour 線條與波形")
        self.btn_clear_contour_lines.setStyleSheet(btn_style_danger)
        self.btn_clear_contour_lines.clicked.connect(self.clear_contour_lines)
        left_layout.addWidget(self.btn_clear_contour_lines)

        left_layout.addWidget(self._create_hline())

        self.lbl_dr_center = QLabel("中心座標: (X: --, Y: --)")
        self.lbl_dr_peak = QLabel("最大強度 (Peak): --")
        self.lbl_dr_thresh_val = QLabel("計算門檻值: --")
        self.lbl_dr_x_width = QLabel("X 軸寬度 (@門檻): -- px (-- μm)")
        self.lbl_dr_y_width = QLabel("Y 軸寬度 (@門檻): -- px (-- μm)")
        self.lbl_dr_area_um = QLabel("實際面積: -- μm²")
        self.lbl_dr_sum_intensity = QLabel("總光強度: --")
        self.lbl_dr_mean_intensity = QLabel("平均強度: --")

        left_layout.addWidget(self.lbl_dr_center)
        left_layout.addWidget(self.lbl_dr_peak)
        left_layout.addWidget(self.lbl_dr_thresh_val)
        left_layout.addWidget(self.lbl_dr_x_width)
        left_layout.addWidget(self.lbl_dr_y_width)
        left_layout.addWidget(self.lbl_dr_area_um)
        left_layout.addWidget(self.lbl_dr_sum_intensity)
        left_layout.addWidget(self.lbl_dr_mean_intensity)
        
        left_layout.addStretch()
        left_scroll_panel.setWidget(left_content)
        left_outer_layout.addWidget(left_scroll_panel)
        
        # --- 右側主繪圖區 ---
        center_panel = QWidget()
        center_layout = QVBoxLayout(center_panel)
        center_layout.setContentsMargins(6, 6, 6, 6)
        center_panel.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
        
        top_bar_layout = QHBoxLayout()
        top_bar_layout.setContentsMargins(0, 0, 0, 6)

        self.btn_select_dir = QPushButton("選擇儲存資料夾")
        self.btn_select_dir.setStyleSheet(btn_style_folder)
        self.btn_select_dir.setMinimumHeight(38)
        self.btn_select_dir.clicked.connect(self.select_save_directory)
        top_bar_layout.addWidget(self.btn_select_dir)

        self.lbl_save_dir_path = QLabel("")
        self.lbl_save_dir_path.setStyleSheet("color: #333333; background-color: #f5f5f5; border: 1px solid #d0d0d0; border-radius: 4px; padding: 6px 10px; font-size: 12px;")
        self.lbl_save_dir_path.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
        self.lbl_save_dir_path.setMinimumHeight(38)
        top_bar_layout.addWidget(self.lbl_save_dir_path)

        self.btn_view_cross_profile = QPushButton("查看十字波形視窗")
        self.btn_view_cross_profile.setStyleSheet(btn_style_cross)
        self.btn_view_cross_profile.setMinimumHeight(38)
        self.btn_view_cross_profile.setEnabled(False)
        self.btn_view_cross_profile.clicked.connect(self.show_cross_profile_window)
        top_bar_layout.addWidget(self.btn_view_cross_profile)

        self.btn_view_m1 = QPushButton("查看 M1 Heatmap")
        self.btn_view_m1.setStyleSheet(btn_style_view)
        self.btn_view_m1.setMinimumHeight(38)
        self.btn_view_m1.setEnabled(False)
        self.btn_view_m1.clicked.connect(self.show_m1_heatmap)
        top_bar_layout.addWidget(self.btn_view_m1)

        self.btn_view_m2 = QPushButton("查看 M2 Heatmap")
        self.btn_view_m2.setStyleSheet(btn_style_view)
        self.btn_view_m2.setMinimumHeight(38)
        self.btn_view_m2.setEnabled(False)
        self.btn_view_m2.clicked.connect(self.show_m2_heatmap)
        top_bar_layout.addWidget(self.btn_view_m2)

        self.btn_export = QPushButton("匯出結果 (JSON/Excel/PNG)")
        self.btn_export.setStyleSheet(btn_style_export)
        self.btn_export.setMinimumHeight(38)
        self.btn_export.setFixedWidth(240)
        self.btn_export.setEnabled(False)
        self.btn_export.clicked.connect(self.export_results)
        top_bar_layout.addWidget(self.btn_export)
        
        center_layout.addLayout(top_bar_layout)
        
        content_splitter = QSplitter(Qt.Horizontal)
        center_layout.addWidget(content_splitter)
        
        left_sub_splitter = QSplitter(Qt.Vertical)
        
        self.win_top = pg.GraphicsLayoutWidget()
        
        colors = [(0, 0, 255), (0, 255, 255), (0, 255, 0), (255, 255, 0), (255, 0, 0)]
        pos = np.linspace(0.0, 1.0, len(colors))
        jet_map = pg.ColorMap(pos, colors)
        
        self.plot_heat = self.win_top.addPlot(row=0, col=0, title='Processed Matrix Result Heatmap')
        self.plot_heat.getViewBox().invertY(False)
        self.plot_heat.setAspectLocked(True)
        self.plot_heat.setLabel('bottom', 'X Pixels')
        self.plot_heat.setLabel('left', 'Y Pixels')
        
        self.image_item = pg.ImageItem()
        self.plot_heat.addItem(self.image_item)

        self.process_thresh_overlay_item = pg.ImageItem()
        self.process_thresh_overlay_item.setZValue(5)
        self.plot_heat.addItem(self.process_thresh_overlay_item)

        self.v_line = pg.InfiniteLine(angle=90, movable=False, pen=pg.mkPen('#37474F', width=1, style=Qt.DashLine))
        self.h_line = pg.InfiniteLine(angle=0, movable=False, pen=pg.mkPen('#37474F', width=1, style=Qt.DashLine))
        self.plot_heat.addItem(self.v_line, ignoreBounds=True)
        self.plot_heat.addItem(self.h_line, ignoreBounds=True)
        
        self.v_line.hide()
        self.h_line.hide()
        
        self.plot_heat.scene().sigMouseMoved.connect(self.mouse_moved)
        self.plot_heat.scene().sigMouseClicked.connect(self.mouse_clicked)
        
        self.hist = LevelAlignedHistogramLUTItem()
        self.hist.setImageItem(self.image_item)
        self.hist.gradient.setColorMap(jet_map)
        self.hist.sigLevelsChanged.connect(self.on_colorbar_levels_changed)
        self.win_top.addItem(self.hist, row=0, col=1)

        apply_readable_plot_theme(self.win_top, [self.plot_heat])
        
        left_sub_splitter.addWidget(self.win_top)
        
        left_scroll = QScrollArea()
        left_scroll.setWidgetResizable(True)
        left_scroll.setStyleSheet("QScrollArea { border: 1px solid #90A4AE; background-color: #E8EEF2; }")
        
        self.win_sub = pg.GraphicsLayoutWidget()
        self.win_sub.setMinimumHeight(400)
        
        self.plot_hist = self.win_sub.addPlot(row=0, col=0, title="Peak Color / Value Distribution")
        self.plot_hist.setLabel('bottom', 'Intensity Value')
        self.plot_hist.setLabel('left', 'Pixel Count')
        self.plot_hist.showGrid(x=True, y=True, alpha=0.3)
        
        self.plot_trend = self.win_sub.addPlot(row=1, col=0, title="Peak Row Trend")
        self.plot_trend.setLabel('bottom', 'X Position (px)')
        self.plot_trend.setLabel('left', 'Intensity')
        self.plot_trend.showGrid(x=True, y=True, alpha=0.3)

        apply_readable_plot_theme(self.win_sub, [self.plot_hist, self.plot_trend])
        
        left_scroll.setWidget(self.win_sub)
        left_sub_splitter.addWidget(left_scroll)
        left_sub_splitter.setStretchFactor(0, 2)
        left_sub_splitter.setStretchFactor(1, 1)
        left_sub_splitter.setSizes([600, 300])
        
        content_splitter.addWidget(left_sub_splitter)
        
        right_sub_splitter = QSplitter(Qt.Vertical)
        
        self.win_contour_top = pg.GraphicsLayoutWidget()
        
        self.plot_contour = self.win_contour_top.addPlot(row=0, col=0, title='Moving Average Contour Map')
        self.plot_contour.getViewBox().invertY(False)
        self.plot_contour.setAspectLocked(True)
        self.plot_contour.setLabel('bottom', 'X Pixels')
        self.plot_contour.setLabel('left', 'Y Pixels')
        
        self.contour_image_item = pg.ImageItem()
        self.plot_contour.addItem(self.contour_image_item)
        
        self.contour_hist = LevelAlignedHistogramLUTItem()
        self.contour_hist.setImageItem(self.contour_image_item)
        self.contour_hist.gradient.setColorMap(jet_map)
        self.contour_hist.sigLevelsChanged.connect(self.on_contour_colorbar_levels_changed)
        self.win_contour_top.addItem(self.contour_hist, row=0, col=1)

        apply_readable_plot_theme(self.win_contour_top, [self.plot_contour])
        
        self.plot_contour.scene().sigMouseClicked.connect(self.on_contour_scene_clicked)
        
        right_sub_splitter.addWidget(self.win_contour_top)
        
        self.win_contour_profile = pg.GraphicsLayoutWidget()
        
        self.plot_contour_profile = self.win_contour_profile.addPlot(title="Contour Profile Waveform (Point 1 -> Point 2)")
        self.plot_contour_profile.setLabel('bottom', 'Distance along line (px)')
        self.plot_contour_profile.setLabel('left', 'Intensity')
        self.plot_contour_profile.showGrid(x=True, y=True, alpha=0.3)
        self.plot_contour_profile.getAxis('left').setWidth(50)

        apply_readable_plot_theme(self.win_contour_profile, [self.plot_contour_profile])
        
        right_sub_splitter.addWidget(self.win_contour_profile)
        right_sub_splitter.setStretchFactor(0, 2)
        right_sub_splitter.setStretchFactor(1, 1)
        right_sub_splitter.setSizes([600, 300])
        
        content_splitter.addWidget(right_sub_splitter)
        content_splitter.setSizes([500, 500])
        
        empty_data = np.zeros((768, 768))
        self.image_item.setImage(empty_data)
        self.contour_image_item.setImage(empty_data)
        
        splitter.addWidget(left_container)
        splitter.addWidget(center_panel)
        splitter.setStretchFactor(0, 0)
        splitter.setStretchFactor(1, 1)
        splitter.setSizes([340, 1060])

    # =========================================================================
    # 事件與邏輯層
    # =========================================================================
    def on_mode_changed(self):
        mode_data = self.combo_mode.currentData()
        is_sub_or_div = (mode_data in ["sub", "div"])
        self.chk_normalize_peaks.setVisible(is_sub_or_div)

        if mode_data == "single":
            self.btn_file1.setText("I. 匯入單一檔案")
            self.btn_file2.setVisible(False)
            self.lbl_file2.setVisible(False)
            self.btn_calc.setText("單獨讀取並畫圖")
            self.lbl_max1.setText("Max: --")
            self.lbl_max2.setVisible(False)
            self.lbl_ratio.setVisible(False)
        elif mode_data == "sub":
            self.btn_file1.setText("I. 匯入第一個檔案")
            self.btn_file2.setText("II. 匯入第二個檔案")
            self.btn_file2.setVisible(True)
            self.lbl_file2.setVisible(True)
            self.btn_calc.setText("雙檔純相減畫圖")
            self.lbl_max1.setText("Max 1: --")
            self.lbl_max2.setText("Max 2: --")
            self.lbl_max2.setVisible(True)
            self.lbl_ratio.setVisible(False)
        elif mode_data == "div":
            self.btn_file1.setText("I. 匯入第一個檔案")
            self.btn_file2.setText("II. 匯入第二個檔案")
            self.btn_file2.setVisible(True)
            self.lbl_file2.setVisible(True)
            self.btn_calc.setText("雙檔純相除畫圖")
            self.lbl_max1.setText("Max 1: --")
            self.lbl_max2.setText("Max 2: --")
            self.lbl_max2.setVisible(True)
            self.lbl_ratio.setVisible(False)
        elif mode_data == "param":
            self.btn_file1.setText("I. 匯入原始 Excel 數據檔")
            self.btn_file2.setText("II. 匯入專屬 JSON 參數檔")
            self.btn_file2.setVisible(True)
            self.lbl_file2.setVisible(True)
            self.btn_calc.setText("載入參數並重繪圖表")
            self.lbl_max1.setText("Max Peak: --")
            self.lbl_max2.setVisible(False)
            self.lbl_ratio.setVisible(False)
        else: # calc
            self.btn_file1.setText("I. 匯入第一個檔案")
            self.btn_file2.setText("II. 匯入第二個檔案")
            self.btn_file2.setVisible(True)
            self.lbl_file2.setVisible(True)
            self.btn_calc.setText("開始計算與畫圖")
            self.lbl_max1.setText("M1最大值(位置): --")
            self.lbl_max2.setText("M2同位置數值: --")
            self.lbl_max2.setVisible(True)
            self.lbl_ratio.setVisible(True)

    def select_save_directory(self):
        dir_path = QFileDialog.getExistingDirectory(self, "選擇儲存資料夾", "")
        if dir_path:
            self.save_dir_path = dir_path
            self.lbl_save_dir_path.setText(f"{dir_path}")

    def show_m1_heatmap(self):
        if self.matrix1 is not None:
            self.apply_m1_point_from_mode()
            if self.viewer_m1_win is not None:
                try:
                    self.viewer_m1_win.close()
                    self.viewer_m1_win.deleteLater()
                except Exception:
                    pass
                self.viewer_m1_win = None

            self.viewer_m1_win = HeatmapViewerWindow("M1 Matrix Heatmap", self.matrix1, app_parent=self, is_m1=True)
            if self.viewer_m2_win and self.viewer_m2_win.isVisible():
                m2_pos = self.viewer_m2_win.pos()
                m1_w = self.viewer_m1_win.width()
                self.viewer_m1_win.setGeometry(max(50, m2_pos.x() - m1_w - 20), m2_pos.y(), 800, 700)
            else:
                self.viewer_m1_win.setGeometry(150, 200, 800, 700)
            self.viewer_m1_win.show()
            self.sync_dual_points_after_m1_change()
            self.update_m1_thresh_overlay()

    def show_m2_heatmap(self):
        if self.matrix2 is not None:
            if self.viewer_m2_win is not None:
                try:
                    self.viewer_m2_win.close()
                    self.viewer_m2_win.deleteLater()
                except Exception:
                    pass
                self.viewer_m2_win = None

            self.viewer_m2_win = HeatmapViewerWindow("M2 Matrix Heatmap", self.matrix2, app_parent=self, is_m1=False)
            if self.viewer_m1_win and self.viewer_m1_win.isVisible():
                m1_pos = self.viewer_m1_win.pos()
                m1_w = self.viewer_m1_win.width()
                self.viewer_m2_win.setGeometry(m1_pos.x() + m1_w + 20, m1_pos.y(), 800, 700)
            else:
                self.viewer_m2_win.setGeometry(980, 200, 800, 700)
            self.viewer_m2_win.show()
            self.update_m2_viewer_markers()
            self.update_m2_thresh_overlay()

    def show_cross_profile_window(self):
        if self.result_matrix is not None:
            if self.cross_profile_win is not None:
                try:
                    self.cross_profile_win.close()
                    self.cross_profile_win.deleteLater()
                except Exception:
                    pass
                self.cross_profile_win = None

            self.cross_profile_win = CrossProfileViewerWindow("DataRay", self)
            self.cross_profile_win.show()
            
            cx, cy = self.heatmap_cross_point if self.heatmap_cross_point is not None else (self.result_matrix.shape[1]//2, self.result_matrix.shape[0]//2)
            min_v, max_v = self.hist.getLevels()
            self.cross_profile_win.update_profiles(self.result_matrix, cx, cy, y_range=(min_v, max_v))

    def update_all_m1_markers(self):
        if self.viewer_m1_win and self.viewer_m1_win.isVisible():
            self.viewer_m1_win.draw_marker(self.m1_center_point)
        self.update_m1_thresh_overlay()
        self.update_m2_viewer_markers()
        if self.result_matrix is not None:
            self.redraw_heatmap_cross_item()
        self.sync_dataray_spot_to_m1()

    def _get_p2_display_point(self):
        if len(self.click_points) >= 2:
            return self.click_points[1]
        return None

    def update_m2_viewer_markers(self):
        if self.viewer_m2_win and self.viewer_m2_win.isVisible():
            self.viewer_m2_win.draw_marker(self.m1_center_point, pt2=self._get_p2_display_point())
        self.update_m2_thresh_overlay()

    def _build_threshold_mask(self, matrix, use_threshold, thresh_percent, y_below=None,
                              robust=True):
        matrix = np.asarray(matrix, dtype=np.float64)
        if matrix.size == 0:
            return None
        h, w = matrix.shape
        if y_below is not None:
            y_below = split_y_index(y_below)
            if y_below <= 0:
                return np.zeros((h, w), dtype=bool)
            region = matrix[:y_below, :]
            region_mask = build_robust_threshold_mask(
                region, use_threshold, thresh_percent,
                bg_subtract=robust, largest_cc_only=robust,
            )
            mask = np.zeros((h, w), dtype=bool)
            if region_mask is not None:
                mask[:y_below, :] = region_mask
            return mask
        return build_robust_threshold_mask(
            matrix, use_threshold, thresh_percent,
            bg_subtract=robust, largest_cc_only=robust,
        )

    def update_m1_thresh_overlay(self, _checked=None):
        if not (self.viewer_m1_win and self.viewer_m1_win.isVisible()):
            return
        show = self.chk_m1_show_thresh.isChecked()
        if (not show) or self.matrix1 is None:
            self.viewer_m1_win.clear_threshold_overlay()
            return
        mask = self._build_threshold_mask(
            self.matrix1,
            self.chk_m1_use_threshold.isChecked(),
            self.spin_m1_thresh_percent.value(),
        )
        self.viewer_m1_win.set_threshold_overlay(mask, visible=True, rgba_color=(255, 64, 255, 90))

    def update_m2_thresh_overlay(self, _checked=None):
        if not (self.viewer_m2_win and self.viewer_m2_win.isVisible()):
            return
        show = self.chk_p2_show_thresh.isChecked()
        if (not show) or self.matrix2 is None or self.m1_center_point is None:
            self.viewer_m2_win.clear_threshold_overlay()
            return
        _x1, y1 = self.m1_center_point
        mask = self._build_threshold_mask(
            self.matrix2,
            self.chk_p2_use_threshold.isChecked(),
            self.spin_p2_thresh_percent.value(),
            y_below=y1,
        )
        self.viewer_m2_win.set_threshold_overlay(mask, visible=True, rgba_color=(255, 64, 255, 90))

    def update_process_thresh_overlay(self, _checked=None):
        if not hasattr(self, "process_thresh_overlay_item") or self.process_thresh_overlay_item is None:
            return
        show = self.chk_dr_show_thresh.isChecked()
        if (not show) or self.result_matrix is None:
            self.process_thresh_overlay_item.clear()
            return
        mask = self._build_threshold_mask(
            self.result_matrix,
            self.chk_dr_use_threshold.isChecked(),
            self.spin_dr_thresh_percent.value(),
            robust=False,
        )
        if mask is None:
            self.process_thresh_overlay_item.clear()
            return
        h, w = mask.shape
        rgba = np.zeros((h, w, 4), dtype=np.ubyte)
        if np.any(mask):
            rgba[mask, 0] = 255
            rgba[mask, 1] = 64
            rgba[mask, 2] = 255
            rgba[mask, 3] = 80
        self.process_thresh_overlay_item.setImage(np.transpose(rgba, (1, 0, 2)), levels=(0, 255))

    def on_dr_enable_spot_toggled(self, checked):
        if checked:
            self.sync_dataray_spot_to_m1()
        else:
            self.clear_dataray_spot_items()

    def sync_dataray_spot_to_m1(self):
        if self.result_matrix is None or not self.chk_dr_enable_spot.isChecked():
            if hasattr(self, "chk_dr_enable_spot") and not self.chk_dr_enable_spot.isChecked():
                self.clear_dataray_spot_items()
            return
        if self.m1_center_point is not None:
            self.dataray_center = tuple(self.m1_center_point)
        else:
            use_thresh = self.chk_dr_use_threshold.isChecked()
            thresh_percent = self.spin_dr_thresh_percent.value()
            self.dataray_center = self._compute_auto_spot_center(
                self.result_matrix, "peak_geom", use_thresh, thresh_percent)
        self.update_dataray_circle()

    def on_m1_point_mode_changed(self, checked=False):
        if checked is False and not any([
            self.radio_m1_peak_geom.isChecked(),
            self.radio_m1_centroid.isChecked(),
            self.radio_m1_thresh_geom.isChecked(),
            self.radio_m1_manual.isChecked(),
        ]):
            return
        sender = self.sender()
        if sender is not None and hasattr(sender, "isChecked") and not sender.isChecked():
            return
        self.apply_m1_point_from_mode()
        self.sync_dual_points_after_m1_change()

    def on_p2_point_mode_changed(self, checked=False):
        sender = self.sender()
        if sender is not None and hasattr(sender, "isChecked") and not sender.isChecked():
            return
        self.apply_p2_point_from_mode()

    def _get_m1_auto_mode_name(self):
        if self.radio_m1_peak_geom.isChecked():
            return "peak_geom"
        if self.radio_m1_centroid.isChecked():
            return "centroid"
        if self.radio_m1_thresh_geom.isChecked():
            return "thresh_geom"
        return "manual"

    def apply_m1_point_from_mode(self):
        if self.matrix1 is None:
            return
        mode = self._get_m1_auto_mode_name()
        if mode == "manual":
            return
        use_thresh = self.chk_m1_use_threshold.isChecked()
        thresh_percent = self.spin_m1_thresh_percent.value()
        self.m1_center_point = self._compute_auto_spot_center(
            self.matrix1, mode, use_thresh, thresh_percent)
        self.update_all_m1_markers()

    def _is_p2_auto_mode(self):
        return (
            self.radio_p2_auto_min.isChecked()
            or self.radio_p2_m2_thresh_geom.isChecked()
            or self.radio_p2_m2_centroid.isChecked()
        )

    def _get_p2_point_mode_name(self):
        if self.radio_p2_auto_min.isChecked():
            return "auto_min"
        if self.radio_p2_m2_thresh_geom.isChecked():
            return "m2_thresh_geom"
        if self.radio_p2_m2_centroid.isChecked():
            return "m2_centroid"
        return "manual"

    def _find_min_below_y(self, matrix, y1):
        matrix = np.asarray(matrix)
        y1 = split_y_index(y1)
        if y1 <= 0:
            return None
        region = matrix[:y1, :]
        if region.size == 0:
            return None
        min_val = np.min(region)
        ys, xs = np.where(region == min_val)
        if len(xs) == 0:
            return None
        cx = float(np.mean(xs))
        cy = float(np.mean(ys))
        return (cx, cy)

    def _find_center_below_y(self, matrix, y1, mode, use_thresh, thresh_percent):
        matrix = np.asarray(matrix)
        y1 = split_y_index(y1)
        if y1 <= 0:
            return None
        region = matrix[:y1, :]
        if region.size == 0:
            return None
        return self._compute_auto_spot_center(region, mode, use_thresh, thresh_percent)

    def apply_p2_point_from_mode(self, silent=False):
        if self.radio_p2_manual.isChecked():
            return
        if self.m1_center_point is None:
            return

        mode = self._get_p2_point_mode_name()
        _x1, y1 = self.m1_center_point

        if mode == "auto_min":
            if self.result_matrix is None:
                return
            if self.matrix1 is not None and self.result_matrix.shape != self.matrix1.shape:
                if not silent:
                    QMessageBox.warning(self, "警告", "M1 與 Process Result 矩陣尺寸不一致，無法自動抓取第二點。")
                return
            p2 = self._find_min_below_y(self.result_matrix, y1)
        else:
            if self.matrix2 is None:
                if not silent:
                    QMessageBox.warning(self, "警告", "尚未載入 M2，無法以此模式自動抓取第二點。")
                return
            if self.matrix1 is not None and self.matrix2.shape != self.matrix1.shape:
                if not silent:
                    QMessageBox.warning(self, "警告", "M1 與 M2 矩陣尺寸不一致，無法自動抓取第二點。")
                return
            use_thresh = self.chk_p2_use_threshold.isChecked()
            thresh_percent = self.spin_p2_thresh_percent.value()
            center_mode = "thresh_geom" if mode == "m2_thresh_geom" else "centroid"
            p2 = self._find_center_below_y(
                self.matrix2, y1, center_mode, use_thresh, thresh_percent)

        if p2 is None:
            if not silent:
                QMessageBox.warning(self, "警告", "第一點 Y 以下沒有可搜尋區域，無法自動抓取第二點。")
            return

        self.click_points = [self.m1_center_point, p2]
        self._ensure_measure_mode_for_dual_points()
        self.update_measure_display()
        self.redraw_measure_crosses()

    def sync_dual_points_after_m1_change(self):
        if self.m1_center_point is None:
            if self._is_p2_auto_mode():
                if len(self.click_points) > 0:
                    self.click_points.clear()
                    self.clear_measure_items_only()
                    self.update_measure_display()
                    self.update_m2_viewer_markers()
            return

        if self._is_p2_auto_mode():
            self.apply_p2_point_from_mode(silent=True)
        else:
            if len(self.click_points) == 0:
                self.click_points = [self.m1_center_point]
            else:
                self.click_points[0] = self.m1_center_point
                if len(self.click_points) > 2:
                    self.click_points = self.click_points[:2]
            self.update_measure_display()
            self.redraw_measure_crosses()

    def _ensure_measure_mode_for_dual_points(self):
        if not self.chk_enable_measure_cross.isChecked():
            self.chk_enable_heatmap_cross.blockSignals(True)
            self.chk_enable_measure_cross.blockSignals(True)
            self.chk_enable_heatmap_cross.setChecked(False)
            self.chk_enable_measure_cross.setChecked(True)
            self.chk_enable_heatmap_cross.blockSignals(False)
            self.chk_enable_measure_cross.blockSignals(False)
            for item in self.heatmap_cross_items:
                self.plot_heat.removeItem(item)
            self.heatmap_cross_items.clear()

    def on_dataray_shape_type_changed(self):
        is_ellipse = self.radio_dr_shape_ellipse.isChecked()
        self.container_dr_circle_spin.setVisible(not is_ellipse)
        self.container_dr_ellipse_spin.setVisible(is_ellipse)
        self.update_dataray_circle()

    def on_dataray_threshold_toggled(self, checked):
        self.spin_dr_thresh_percent.setEnabled(checked)
        self.lbl_dr_thresh_spin.setEnabled(checked)
        self.recalculate_dataray_spot()
        self.update_process_thresh_overlay()

    def on_dr_thresh_percent_changed(self, _value=None):
        self.recalculate_dataray_spot()
        self.update_process_thresh_overlay()

    def on_m1_threshold_toggled(self, checked):
        self.spin_m1_thresh_percent.setEnabled(checked)
        self.lbl_m1_thresh_spin.setEnabled(checked)
        if self.radio_m1_thresh_geom.isChecked() or self.radio_m1_centroid.isChecked():
            self.apply_m1_point_from_mode()
            self.sync_dual_points_after_m1_change()
        self.update_m1_thresh_overlay()

    def on_m1_thresh_percent_changed(self, _value=None):
        if self.radio_m1_thresh_geom.isChecked() or self.radio_m1_centroid.isChecked():
            self.apply_m1_point_from_mode()
            self.sync_dual_points_after_m1_change()
        self.update_m1_thresh_overlay()

    def on_p2_threshold_toggled(self, checked):
        self.spin_p2_thresh_percent.setEnabled(checked)
        self.lbl_p2_thresh_spin.setEnabled(checked)
        if self.radio_p2_m2_thresh_geom.isChecked() or self.radio_p2_m2_centroid.isChecked():
            self.apply_p2_point_from_mode(silent=True)
        self.update_m2_thresh_overlay()

    def on_p2_thresh_percent_changed(self, _value=None):
        if self.radio_p2_m2_thresh_geom.isChecked() or self.radio_p2_m2_centroid.isChecked():
            self.apply_p2_point_from_mode(silent=True)
        self.update_m2_thresh_overlay()

    def _compute_auto_spot_center(self, matrix, mode, use_threshold=False, thresh_percent=50.0):
        # 強化定位：背景扣除 + 最大連通區 + 亞像素（peak_geom 仍走峰值幾何）
        return compute_auto_spot_center(
            matrix, mode, use_threshold, thresh_percent,
            bg_subtract=(mode != "peak_geom"),
            largest_cc_only=(mode != "peak_geom"),
            subpixel=True,
        )

    def recalculate_dataray_spot(self):
        self.sync_dataray_spot_to_m1()
        self.update_process_thresh_overlay()

    def clear_dataray_spot_items(self):
        if self.dataray_circle_item is not None:
            self.plot_heat.removeItem(self.dataray_circle_item)
            self.dataray_circle_item = None
        if self.dataray_center_spot is not None:
            self.plot_heat.removeItem(self.dataray_center_spot)
            self.dataray_center_spot = None
            
        for item in self.dataray_fixed_cross_items:
            self.plot_heat.removeItem(item)
        self.dataray_fixed_cross_items.clear()

        self.lbl_dr_center.setText("中心座標: (X: --, Y: --)")
        self.lbl_dr_peak.setText("最大強度 (Peak): --")
        self.lbl_dr_thresh_val.setText("計算門檻值: 未啟用" if not self.chk_dr_use_threshold.isChecked() else "計算門檻值: --")
        self.lbl_dr_x_width.setText("X 軸寬度 (@門檻): --")
        self.lbl_dr_y_width.setText("Y 軸寬度 (@門檻): --")
        self.lbl_dr_area_um.setText("實際面積: -- μm²")
        self.lbl_dr_sum_intensity.setText("總光強度: --")
        self.lbl_dr_mean_intensity.setText("平均強度: --")

    def update_dataray_circle(self):
        if self.result_matrix is None or not self.chk_dr_enable_spot.isChecked() or self.dataray_center is None or self.is_updating_dataray_ui:
            if hasattr(self, "chk_dr_enable_spot") and not self.chk_dr_enable_spot.isChecked():
                self.clear_dataray_spot_items()
            return
            
        self.is_updating_dataray_ui = True
        self.clear_dataray_spot_items()

        cx, cy = self.dataray_center
        h, w = self.result_matrix.shape

        peak_value = np.max(self.result_matrix)
        use_thresh = self.chk_dr_use_threshold.isChecked()
        is_ellipse = self.radio_dr_shape_ellipse.isChecked()
        pixel_pitch_um = 5.5

        cy_clamped = max(0, min(h - 1, int(round(cy))))
        x_profile = self.result_matrix[cy_clamped, :]

        cx_clamped = max(0, min(w - 1, int(round(cx))))
        y_profile = self.result_matrix[:, cx_clamped]

        x_width_px, y_width_px = 0, 0
        x_width_um, y_width_um = 0.0, 0.0

        if use_thresh:
            thresh_percent = self.spin_dr_thresh_percent.value()
            thresh_val = peak_value * (thresh_percent / 100.0)

            x_above = np.where(x_profile >= thresh_val)[0]
            if len(x_above) > 1:
                x_width_px = x_above[-1] - x_above[0]
                x_width_um = x_width_px * pixel_pitch_um

            y_above = np.where(y_profile >= thresh_val)[0]
            if len(y_above) > 1:
                y_width_px = y_above[-1] - y_above[0]
                y_width_um = y_width_px * pixel_pitch_um

            if is_ellipse:
                if x_width_px > 0:
                    self.spin_dr_ellipse_wx.blockSignals(True)
                    self.spin_dr_ellipse_wx.setValue(int(x_width_px))
                    self.spin_dr_ellipse_wx.blockSignals(False)
                if y_width_px > 0:
                    self.spin_dr_ellipse_wy.blockSignals(True)
                    self.spin_dr_ellipse_wy.setValue(int(y_width_px))
                    self.spin_dr_ellipse_wy.blockSignals(False)
            else:
                auto_diameter_px = max(x_width_px, y_width_px)
                if auto_diameter_px > 0:
                    self.spin_dr_circle_diameter.blockSignals(True)
                    self.spin_dr_circle_diameter.setValue(int(auto_diameter_px))
                    self.spin_dr_circle_diameter.blockSignals(False)

            self.lbl_dr_thresh_val.setText(f"計算門檻值 ({thresh_percent}%): {thresh_val:.1f}")
            self.lbl_dr_x_width.setText(f"X 軸寬度 (@門檻): {x_width_px} px ({x_width_um:.2f} μm)")
            self.lbl_dr_y_width.setText(f"Y 軸寬度 (@門檻): {y_width_px} px ({y_width_um:.2f} μm)")
        else:
            self.lbl_dr_thresh_val.setText("計算門檻值: 未啟用")
            self.lbl_dr_x_width.setText("X 軸寬度 (@門檻): --")
            self.lbl_dr_y_width.setText("Y 軸寬度 (@門檻): --")

        theta = np.linspace(0, 2*np.pi, 100)

        if is_ellipse:
            wx_px = self.spin_dr_ellipse_wx.value()
            wy_px = self.spin_dr_ellipse_wy.value()
            rx_px = wx_px / 2.0
            ry_px = wy_px / 2.0

            circle_x = cx + rx_px * np.cos(theta)
            circle_y = cy + ry_px * np.sin(theta)

            rx_um = rx_px * pixel_pitch_um
            ry_um = ry_px * pixel_pitch_um
            area_um2 = np.pi * rx_um * ry_um

            y_grid, x_grid = np.ogrid[:h, :w]
            mask = ((x_grid - cx)**2 / (rx_px**2)) + ((y_grid - cy)**2 / (ry_px**2)) <= 1.0
        else:
            diameter_px = self.spin_dr_circle_diameter.value()
            radius_px = diameter_px / 2.0

            circle_x = cx + radius_px * np.cos(theta)
            circle_y = cy + radius_px * np.sin(theta)

            radius_um = radius_px * pixel_pitch_um
            area_um2 = np.pi * (radius_um ** 2)

            y_grid, x_grid = np.ogrid[:h, :w]
            mask = (x_grid - cx)**2 + (y_grid - cy)**2 <= radius_px**2

        v_fixed = pg.PlotCurveItem(x=[cx, cx], y=[0, h], pen=pg.mkPen('y', width=1, style=Qt.DashLine))
        h_fixed = pg.PlotCurveItem(x=[0, w], y=[cy, cy], pen=pg.mkPen('y', width=1, style=Qt.DashLine))
        self.plot_heat.addItem(v_fixed)
        self.plot_heat.addItem(h_fixed)
        self.dataray_fixed_cross_items.extend([v_fixed, h_fixed])

        self.dataray_center_spot = pg.ScatterPlotItem(x=[cx], y=[cy], symbol='+', size=12, pen=pg.mkPen('r', width=2))
        self.plot_heat.addItem(self.dataray_center_spot)
        
        self.dataray_circle_item = pg.PlotCurveItem(circle_x, circle_y, pen=pg.mkPen('y', width=2))
        self.plot_heat.addItem(self.dataray_circle_item)

        circle_pixels = self.result_matrix[mask]
        sum_intensity = np.sum(circle_pixels) if len(circle_pixels) > 0 else 0
        mean_intensity = np.mean(circle_pixels) if len(circle_pixels) > 0 else 0

        self.lbl_dr_center.setText(f"中心座標: (X: {cx:.2f}, Y: {cy:.2f})")
        self.lbl_dr_peak.setText(f"最大強度 (Peak): {peak_value:.1f}")
        self.lbl_dr_area_um.setText(f"實際面積: {area_um2:.2f} μm²")
        self.lbl_dr_sum_intensity.setText(f"總光強度: {sum_intensity:.1f}")
        self.lbl_dr_mean_intensity.setText(f"平均強度: {mean_intensity:.2f}")

        self.is_updating_dataray_ui = False

    def update_measure_display(self):
        pixel_pitch_um = 5.5
        if len(self.click_points) == 1:
            p1 = self.click_points[0]
            self.lbl_distance.setText("位置差距: ΔX: --, ΔY: -- \n總距離: -- px")
            self.lbl_real_distance.setText("實際差距 (* 5.5): ΔX: -- μm, ΔY: -- μm \n總距離: -- μm")
        elif len(self.click_points) == 2:
            p1 = self.click_points[0]
            p2 = self.click_points[1]
            dx = p2[0] - p1[0]
            dy = p2[1] - p1[1]
            distance_px = np.sqrt(dx**2 + dy**2)
            self.lbl_distance.setText(
                f"位置差距: ΔX: {abs(dx):.2f} px, ΔY: {abs(dy):.2f} px \n總距離: {distance_px:.2f} px"
            )
            dx_real = abs(dx) * pixel_pitch_um
            dy_real = abs(dy) * pixel_pitch_um
            distance_real = distance_px * pixel_pitch_um
            self.lbl_real_distance.setText(
                f"實際差距 (* 5.5): ΔX: {dx_real:.2f} μm, ΔY: {dy_real:.2f} μm \n總距離: {distance_real:.2f} μm"
            )

    def clear_measure_points(self):
        self.click_points.clear()
        self.clear_measure_items_only()
        self.lbl_distance.setText("位置差距: ΔX: --, ΔY: -- \n總距離: -- px")
        self.lbl_real_distance.setText("實際差距 (* 5.5): ΔX: -- μm, ΔY: -- μm \n總距離: -- μm")

    def mouse_moved(self, evt):
        pos = evt
        if self.plot_heat.sceneBoundingRect().contains(pos):
            mouse_point = self.plot_heat.getViewBox().mapSceneToView(pos)
            x = mouse_point.x()
            y = mouse_point.y()
            if self.result_matrix is not None:
                h, w = self.result_matrix.shape
                if 0 <= x < w and 0 <= y < h:
                    self.v_line.show()
                    self.h_line.show()
                    self.v_line.setPos(x)
                    self.h_line.setPos(y)
                    self.lbl_cursor.setText(f"目前滑鼠: X: {x:.1f}, Y: {y:.1f}")
                else:
                    self.v_line.hide()
                    self.h_line.hide()
                    self.lbl_cursor.setText("目前滑鼠: X: --, Y: --")
            else:
                self.v_line.hide()
                self.h_line.hide()
        else:
            self.v_line.hide()
            self.h_line.hide()
            self.lbl_cursor.setText("目前滑鼠: X: --, Y: --")

    def mouse_clicked(self, evt):
        if self.result_matrix is None:
            return
        pos = evt.scenePos()
        if self.plot_heat.sceneBoundingRect().contains(pos):
            mouse_point = self.plot_heat.getViewBox().mapSceneToView(pos)
            x = int(round(mouse_point.x()))
            y = int(round(mouse_point.y()))
            h, w = self.result_matrix.shape
            if 0 <= x < w and 0 <= y < h:
                if self.chk_enable_heatmap_cross.isChecked():
                    self.heatmap_cross_point = (x, y)
                    self.redraw_heatmap_cross_item()

                    if self.cross_profile_win is not None and self.cross_profile_win.isVisible():
                        min_v, max_v = self.hist.getLevels()
                        self.cross_profile_win.update_profiles(self.result_matrix, x, y, y_range=(min_v, max_v))

                if self.radio_p2_manual.isChecked():
                    if not self.chk_enable_measure_cross.isChecked():
                        self._ensure_measure_mode_for_dual_points()

                    if self.m1_center_point is not None and len(self.click_points) == 0:
                        self.click_points.append(self.m1_center_point)

                    if len(self.click_points) >= 2:
                        self.clear_measure_items_only()
                        self.click_points.clear()
                        if self.m1_center_point is not None:
                            self.click_points.append(self.m1_center_point)

                    if (x, y) not in self.click_points:
                        self.click_points.append((x, y))

                    self.update_measure_display()
                    self.redraw_measure_crosses()
                elif self.chk_enable_measure_cross.isChecked() and not self._is_p2_auto_mode():
                    if self.m1_center_point is not None and len(self.click_points) == 0:
                        self.click_points.append(self.m1_center_point)

                    if len(self.click_points) >= 2:
                        self.clear_measure_items_only()
                        self.click_points.clear()
                        if self.m1_center_point is not None:
                            self.click_points.append(self.m1_center_point)

                    if (x, y) not in self.click_points:
                        self.click_points.append((x, y))

                    self.update_measure_display()
                    self.redraw_measure_crosses()

    def redraw_heatmap_cross_item(self):
        for item in self.heatmap_cross_items:
            self.plot_heat.removeItem(item)
        self.heatmap_cross_items.clear()

        if self.heatmap_cross_point is not None and self.chk_enable_heatmap_cross.isChecked():
            cx, cy = self.heatmap_cross_point
            h, w = self.result_matrix.shape
            pen = pg.mkPen('#00E676', width=2)
            v_item = pg.PlotCurveItem(x=[cx, cx], y=[0, h], pen=pen)
            h_item = pg.PlotCurveItem(x=[0, w], y=[cy, cy], pen=pen)
            self.plot_heat.addItem(v_item)
            self.plot_heat.addItem(h_item)
            self.heatmap_cross_items.extend([v_item, h_item])

        for item in self.m1_marker_items:
            self.plot_heat.removeItem(item)
        self.m1_marker_items.clear()

        if self.m1_center_point is not None:
            mcx, mcy = self.m1_center_point
            h, w = self.result_matrix.shape
            m_pen = pg.mkPen('#76FF03', width=1.5, style=Qt.DashLine)
            m_v = pg.PlotCurveItem(x=[mcx, mcx], y=[0, h], pen=m_pen)
            m_h = pg.PlotCurveItem(x=[0, w], y=[mcy, mcy], pen=m_pen)
            self.plot_heat.addItem(m_v)
            self.plot_heat.addItem(m_h)
            self.m1_marker_items.extend([m_v, m_h])

    def redraw_measure_crosses(self):
        self.clear_measure_items_only()
        if not self.chk_enable_measure_cross.isChecked():
            self.update_m2_viewer_markers()
            return
        half_size = self.spin_cross_size.value() / 2.0
        for idx, (cx, cy) in enumerate(self.click_points):
            color = 'y' if idx == 0 else 'c'
            pen = pg.mkPen(color, width=2)
            v_marker = pg.PlotCurveItem(x=[cx, cx], y=[cy - half_size, cy + half_size], pen=pen)
            h_marker = pg.PlotCurveItem(x=[cx - half_size, cx + half_size], y=[cy, cy], pen=pen)
            self.plot_heat.addItem(v_marker)
            self.plot_heat.addItem(h_marker)
            self.measure_items.extend([v_marker, h_marker])
        self.update_m2_viewer_markers()

    def clear_measure_items_only(self):
        for item in self.measure_items:
            self.plot_heat.removeItem(item)
        self.measure_items.clear()

    def on_heatmap_cross_toggled(self, checked):
        if checked:
            self.chk_enable_measure_cross.blockSignals(True)
            self.chk_enable_measure_cross.setChecked(False)
            self.chk_enable_measure_cross.blockSignals(False)
            self.clear_measure_items_only()
        self.redraw_heatmap_cross_item()

    def on_measure_cross_toggled(self, checked):
        if checked:
            self.chk_enable_heatmap_cross.blockSignals(True)
            self.chk_enable_heatmap_cross.setChecked(False)
            self.chk_enable_heatmap_cross.blockSignals(False)
            for item in self.heatmap_cross_items:
                self.plot_heat.removeItem(item)
            self.heatmap_cross_items.clear()
            self.sync_dual_points_after_m1_change()
        self.redraw_measure_crosses()

    def clear_sub_plots(self):
        self.plot_hist.clear()
        self.plot_trend.clear()

    def render_sub_plots_fast(self, matrix, peak_row):
        self.clear_sub_plots()
        total_pixels = matrix.size
        sample_data = matrix.ravel()[::10] if total_pixels > 1000000 else matrix.ravel()
        y_counts, x_edges = np.histogram(sample_data, bins=40)
        
        bar_item = pg.BarGraphItem(x0=x_edges[:-1], x1=x_edges[1:], height=y_counts, brush='#E91E63', pen=None)
        self.plot_hist.addItem(bar_item)
        
        peak_val = np.max(matrix)
        v_line_peak = pg.InfiniteLine(pos=peak_val, angle=90, pen=pg.mkPen('r', width=2, style=Qt.DashLine))
        self.plot_hist.addItem(v_line_peak)
        
        min_lvl, max_lvl = self.hist.getLevels()
        self.plot_hist.setXRange(min_lvl, max_lvl, padding=0)
        
        row_profile = matrix[peak_row, :]
        x_axis = np.arange(len(row_profile))
        
        trend_curve = pg.PlotCurveItem(x_axis, row_profile, pen=pg.mkPen('#00E5FF', width=1.5))
        peak_col = np.argmax(row_profile)
        peak_spot = pg.ScatterPlotItem(x=[peak_col], y=[row_profile[peak_col]], symbol='o', size=8, brush='y', pen='r')
        
        self.plot_trend.addItem(trend_curve)
        self.plot_trend.addItem(peak_spot)
        self.plot_trend.setTitle(f"Peak Row Trend (Row Index: {peak_row})")

    def on_contour_scene_clicked(self, evt):
        if self.smoothed_matrix is None:
            return
        pos = evt.scenePos()
        if self.plot_contour.sceneBoundingRect().contains(pos):
            mouse_point = self.plot_contour.getViewBox().mapSceneToView(pos)
            x = mouse_point.x()
            y = mouse_point.y()
            h, w = self.smoothed_matrix.shape
            if 0 <= x < w and 0 <= y < h:
                if len(self.contour_click_points) >= 2:
                    self.clear_contour_measure_items()
                    self.contour_click_points.clear()
                self.contour_click_points.append((x, y))
                self.redraw_contour_overlay()

    def clear_contour_measure_items(self):
        for item in self.contour_measure_items:
            self.plot_contour.removeItem(item)
        self.contour_measure_items.clear()

    def redraw_contour_overlay(self):
        self.clear_contour_measure_items()
        if len(self.contour_click_points) == 1:
            p1 = self.contour_click_points[0]
            spot = pg.ScatterPlotItem(x=[p1[0]], y=[p1[1]], symbol='o', size=8, brush='r', pen='w')
            self.plot_contour.addItem(spot)
            self.contour_measure_items.append(spot)
            self.plot_contour_profile.clear()
        elif len(self.contour_click_points) == 2:
            p1, p2 = self.contour_click_points[0], self.contour_click_points[1]
            line = pg.PlotCurveItem(x=[p1[0], p2[0]], y=[p1[1], p2[1]], pen=pg.mkPen('r', width=2))
            spots = pg.ScatterPlotItem(x=[p1[0], p2[0]], y=[p1[1], p2[1]], symbol='o', size=8, brush='r', pen='w')
            self.plot_contour.addItem(line)
            self.plot_contour.addItem(spots)
            self.contour_measure_items.extend([line, spots])
            
            x_axis_data, profile, x_label = self.calculate_line_profile(self.smoothed_matrix, p1, p2)
            if profile is not None:
                self.plot_contour_profile.clear()
                c_min, c_max = self.contour_hist.getLevels()
                clipped_profile = np.clip(profile, c_min, c_max)
                
                curve = pg.PlotCurveItem(x_axis_data, clipped_profile, pen=pg.mkPen('#FF5722', width=2))
                self.plot_contour_profile.addItem(curve)
                self.plot_contour_profile.setLabel('bottom', x_label)
                self.plot_contour_profile.setXRange(min(x_axis_data), max(x_axis_data), padding=0)
                self.plot_contour_profile.setYRange(c_min, c_max)
                self.plot_contour_profile.setTitle(f"Contour Profile Waveform: ({p1[0]:.0f},{p1[1]:.0f}) -> ({p2[0]:.0f},{p2[1]:.0f})")

    def calculate_line_profile(self, matrix, p1, p2):
        x0, y0 = p1
        x1, y1 = p2
        length = int(np.hypot(x1 - x0, y1 - y0))
        if length < 2:
            return None, None, ""
        x_coords = np.linspace(x0, x1, length)
        y_coords = np.linspace(y0, y1, length)
        profile = map_coordinates(matrix, [y_coords, x_coords], order=1)
        
        if abs(x1 - x0) >= abs(y1 - y0):
            x_axis_data, x_label = x_coords, "X Position (px)"
        else:
            x_axis_data, x_label = y_coords, "Y Position (px)"
            
        return x_axis_data, profile, x_label

    def render_contour_map(self, matrix):
        k_size = self.spin_ma_size.value()
        self.smoothed_matrix = uniform_filter(matrix, size=k_size, mode='nearest') if k_size > 1 else matrix
        self.contour_image_item.setImage(self.smoothed_matrix.T)
        self.plot_contour.setTitle(f'Moving Average Contour (Kernel: {k_size}x{k_size})')
        
        for c_item in self.contour_curves:
            self.plot_contour.removeItem(c_item)
        self.contour_curves.clear()
        
        min_v, max_v = float(np.min(self.smoothed_matrix)), float(np.max(self.smoothed_matrix))
        self.contour_hist.setHistogramRange(min_v, max_v, padding=0)
        self.contour_hist.setLevels(min_v, max_v)
        self.update_isocurves_and_waveform()

    def load_file1(self):
        path, _ = QFileDialog.getOpenFileName(
            self,
            "選擇數據檔案",
            "",
            "Data Files (*.xlsx *.xls *.csv *.npy)"
        )
        if path:
            self.file1_path = path
            self.lbl_file1.setText(os.path.basename(path))
            self.lbl_file1.setStyleSheet("color: #212121; font-size: 11px;")
            if not self.save_dir_path:
                self.save_dir_path = os.path.dirname(path)
            mode_data = self.combo_mode.currentData()
            if mode_data not in ["single", "param"]:
                QTimer.singleShot(100, self.load_file2)

    def load_file2(self):
        mode_data = self.combo_mode.currentData()
        if mode_data == "param":
            path, _ = QFileDialog.getOpenFileName(self, "選擇 JSON 參數檔案", "", "JSON Files (*.json)")
            if path:
                self.param_file_path = path
                self.lbl_file2.setText(os.path.basename(path))
                self.lbl_file2.setStyleSheet("color: #212121; font-size: 11px;")
        else:
            path, _ = QFileDialog.getOpenFileName(
                self,
                "選擇第二個數據檔案",
                "",
                "Data Files (*.xlsx *.xls *.csv *.npy)"
            )
            if path:
                self.file2_path = path
                self.lbl_file2.setText(os.path.basename(path))
                self.lbl_file2.setStyleSheet("color: #212121; font-size: 11px;")

    def process_data(self):
        mode_data = self.combo_mode.currentData()
        is_dual_mode = (mode_data not in ["single", "param"])
        if is_dual_mode and (not self.file1_path or not self.file2_path):
            QMessageBox.warning(self, "警告", "目前為雙檔運算模式，請務必完整匯入第一個與第二個檔案！")
            return

        self.clear_measure_points() 
        self.contour_click_points.clear()
        if mode_data == "single":
            self.plot_single_file()
        elif mode_data == "sub":
            self.plot_pure_subtraction()
        elif mode_data == "div":
            self.plot_pure_division()
        elif mode_data == "param":
            self.plot_with_parameters()
        else:
            self.calculate_and_plot()

    def export_results(self):
        if self.result_matrix is None:
            QMessageBox.warning(self, "警告", "目前無可匯出的數據！")
            return
            
        if not self.save_dir_path:
            QMessageBox.warning(self, "警告", "請先點擊「選擇儲存資料夾」按鈕以指定儲存路徑！")
            return

        initial_path = "Result"
        if self.save_dir_path:
            initial_path = os.path.join(self.save_dir_path, "Result")

        path, _ = QFileDialog.getSaveFileName(self, "選擇儲存主檔名", initial_path, "JSON Files (*.json)")
        if not path:
            return

        base_path, ext = os.path.splitext(path)
        if ext.lower() != ".json":
            path = base_path + ".json"

        try:
            self.lbl_status.setText("狀態: 正在匯出檔案...")
            self.lbl_status.setStyleSheet("color: #F57C00; font-weight: bold;")
            self.btn_export.setEnabled(False)
            QApplication.processEvents()

            c_min, c_max = self.contour_hist.getLevels()
            cx, cy = self.dataray_center if self.dataray_center else (0, 0)
            is_ellipse = self.radio_dr_shape_ellipse.isChecked()

            # 將設定與定位功能及門檻比例完整寫入 JSON 參數檔
            config_params = {
                "ma_kernel_size": self.spin_ma_size.value(),
                "colorbar_min": c_min,
                "colorbar_max": c_max,
                "cross_size": self.spin_cross_size.value(),
                "heatmap_click_points": self.click_points,
                "contour_click_points": self.contour_click_points,
                "dr_enable_spot": self.chk_dr_enable_spot.isChecked(),
                "m1_center_x_px": self.m1_center_point[0] if self.m1_center_point else None,
                "m1_center_y_px": self.m1_center_point[1] if self.m1_center_point else None,
                "m1_point_mode": self._get_m1_auto_mode_name(),
                "p2_point_mode": self._get_p2_point_mode_name(),
                "dr_shape_type": "ellipse" if is_ellipse else "circle",
                "dr_center_x_px": cx,
                "dr_center_y_px": cy,
                "dr_use_threshold": self.chk_dr_use_threshold.isChecked(),
                "dr_threshold_percent": self.spin_dr_thresh_percent.value(),
                "m1_use_threshold": self.chk_m1_use_threshold.isChecked(),
                "m1_threshold_percent": self.spin_m1_thresh_percent.value(),
                "p2_use_threshold": self.chk_p2_use_threshold.isChecked(),
                "p2_threshold_percent": self.spin_p2_thresh_percent.value(),
                "dr_circle_diameter_px": self.spin_dr_circle_diameter.value(),
                "dr_ellipse_wx_px": self.spin_dr_ellipse_wx.value(),
                "dr_ellipse_wy_px": self.spin_dr_ellipse_wy.value()
            }

            with open(path, 'w', encoding='utf-8') as f:
                json.dump(config_params, f, indent=4, ensure_ascii=False)

            mode_data = self.combo_mode.currentData()
            is_dual_file_mode = (mode_data in ["calc", "div", "sub"])
            excel_saved_msg = ""
            if is_dual_file_mode:
                excel_img_path = f"{base_path}_Result.xlsx"
                pd.DataFrame(self.result_matrix).to_excel(excel_img_path, index=False, header=False)
                excel_saved_msg = f"算完數據 Excel: {os.path.basename(excel_img_path)}\n"

                csv_img_path = f"{base_path}_Result.csv"
                pd.DataFrame(self.result_matrix).to_csv(csv_img_path, index=False, header=False)
                excel_saved_msg += f"算完數據 CSV: {os.path.basename(csv_img_path)}\n"

                npy_img_path = f"{base_path}_Result.npy"
                np.save(npy_img_path, self.result_matrix)
                excel_saved_msg += f"算完數據 NPY: {os.path.basename(npy_img_path)}\n"

            spot_analysis_excel_path = f"{base_path}_Spot_Analysis.xlsx"
            wb_spot = openpyxl.Workbook()
            ws_spot = wb_spot.active
            ws_spot.title = "Spot_and_Measurement"
            ws_spot.append(["Item", "Value", "Unit"])

            p1 = self.click_points[0] if len(self.click_points) > 0 else ("--", "--")
            p2 = self.click_points[1] if len(self.click_points) > 1 else ("--", "--")
            
            dx_px, dy_px, dist_px = "--", "--", "--"
            dx_um, dy_um, dist_um = "--", "--", "--"
            pixel_pitch_um = 5.5

            if len(self.click_points) == 2:
                dx_px = abs(self.click_points[1][0] - self.click_points[0][0])
                dy_px = abs(self.click_points[1][1] - self.click_points[0][1])
                dist_px = np.sqrt(dx_px**2 + dy_px**2)
                dx_um = dx_px * pixel_pitch_um
                dy_um = dy_px * pixel_pitch_um
                dist_um = dist_px * pixel_pitch_um

            spot_rows = [
                ["Shape Type", "Ellipse" if is_ellipse else "Circle", ""],
                ["Circle Size / Wx", self.spin_dr_ellipse_wx.value() if is_ellipse else self.spin_dr_circle_diameter.value(), "px"],
                ["Ellipse Wy", self.spin_dr_ellipse_wy.value() if is_ellipse else "N/A", "px"],
                ["Use Spot Threshold", "Yes" if self.chk_dr_use_threshold.isChecked() else "No", ""],
                ["Spot Threshold Percent", self.spin_dr_thresh_percent.value(), "%"],
                ["M1 Point Mode", self._get_m1_auto_mode_name(), ""],
                ["Use M1 Point Threshold", "Yes" if self.chk_m1_use_threshold.isChecked() else "No", ""],
                ["M1 Point Threshold Percent", self.spin_m1_thresh_percent.value(), "%"],
                ["P2 Point Mode", self._get_p2_point_mode_name(), ""],
                ["Use P2 Point Threshold", "Yes" if self.chk_p2_use_threshold.isChecked() else "No", ""],
                ["P2 Point Threshold Percent", self.spin_p2_thresh_percent.value(), "%"],
                ["Mouse Cursor X", "N/A (Realtime)", "px"],
                ["Mouse Cursor Y", "N/A (Realtime)", "px"],
                ["Click Point 1 (X)", p1[0] if p1 != ("--", "--") else "--", "px"],
                ["Click Point 1 (Y)", p1[1] if p1 != ("--", "--") else "--", "px"],
                ["Click Point 2 (X)", p2[0] if p2 != ("--", "--") else "--", "px"],
                ["Click Point 2 (Y)", p2[1] if p2 != ("--", "--") else "--", "px"],
                ["Delta X (px)", dx_px, "px"],
                ["Delta Y (px)", dy_px, "px"],
                ["Total Distance (px)", dist_px, "px"],
                ["Delta X (Real)", dx_um, "μm"],
                ["Delta Y (Real)", dy_um, "μm"],
                ["Total Distance (Real)", dist_um, "μm"],
                ["Cross Marker Size", self.spin_cross_size.value(), "px"]
            ]
            for r in spot_rows:
                ws_spot.append(r)
            wb_spot.save(spot_analysis_excel_path)
            spot_analysis_msg = f"光斑與量測數據 Excel: {os.path.basename(spot_analysis_excel_path)}\n"

            contour_img_path = f"{base_path}_Contour{EXPORT_IMAGE_EXT}"
            export_plot_image(self.plot_contour, contour_img_path)

            heatmap_img_path = f"{base_path}_Heatmap{EXPORT_IMAGE_EXT}"
            export_plot_image(self.plot_heat, heatmap_img_path)

            waveform_excel_msg = ""
            if len(self.contour_click_points) == 2 and self.smoothed_matrix is not None:
                p1_c, p2_c = self.contour_click_points[0], self.contour_click_points[1]
                x_axis_data, profile, x_label = self.calculate_line_profile(self.smoothed_matrix, p1_c, p2_c)
                
                if profile is not None:
                    waveform_excel_path = f"{base_path}_Waveform.xlsx"
                    clipped_profile = np.clip(profile, c_min, c_max)
                    
                    wb = openpyxl.Workbook()
                    ws = wb.active
                    ws.title = "Waveform Profile Data"
                    ws.append([x_label, "Intensity"])
                    
                    for x_val, y_val in zip(x_axis_data, clipped_profile):
                        ws.append([float(x_val), float(y_val)])
                        
                    chart = LineChart()
                    chart.title = f"Contour Waveform Profile ({p1_c[0]:.0f},{p1_c[1]:.0f}) -> ({p2_c[0]:.0f},{p2_c[1]:.0f})"
                    chart.style = 13
                    chart.y_axis.title = "Intensity"
                    chart.x_axis.title = x_label
                    
                    data_ref = Reference(ws, min_col=2, min_row=1, max_row=len(clipped_profile)+1)
                    cats_ref = Reference(ws, min_col=1, min_row=2, max_row=len(clipped_profile)+1)
                    chart.add_data(data_ref, titles_from_data=True)
                    chart.set_categories(cats_ref)
                    chart.width, chart.height = 16, 10
                    
                    ws.add_chart(chart, "D2")
                    wb.save(waveform_excel_path)
                    waveform_excel_msg = f"Waveform 數據與圖表 Excel: {os.path.basename(waveform_excel_path)}\n"

            self.lbl_status.setText("狀態: 所有檔案匯出成功！")
            self.lbl_status.setStyleSheet("color: #2E7D32; font-weight: bold;")
            self.btn_export.setEnabled(True)

            QMessageBox.information(
                self, "成功", 
                f"匯出完成！已儲存至：\n\n{result_export_msg}"
                f"{spot_analysis_msg}"
                f"JSON 參數檔: {os.path.basename(path)}\n"
                f"Contour 圖: {os.path.basename(contour_img_path)}\n"
                f"Heatmap 圖: {os.path.basename(heatmap_img_path)}\n"
                f"{waveform_excel_msg}"
            )
        except Exception as e:
            self.lbl_status.setText("狀態: 匯出失敗")
            self.lbl_status.setStyleSheet("color: #C62828; font-weight: bold;")
            self.btn_export.setEnabled(True)
            QMessageBox.critical(self, "匯出錯誤", f"匯出過程發生錯誤：\n{str(e)}")

    def plot_with_parameters(self):
        if not self.file1_path or not self.param_file_path:
            QMessageBox.warning(self, "警告", "請確保數據檔與 JSON 參數檔皆已選取！")
            return
        try:
            self.lbl_status.setText("狀態: 正在讀取數據與解析參數檔...")
            self.lbl_status.setStyleSheet("color: #F57C00; font-weight: bold;")
            QApplication.processEvents()

            self.matrix1 = load_numeric_matrix(self.file1_path)
            self.matrix2 = None
            self.result_matrix = self.matrix1

            self.btn_view_m1.setEnabled(True)
            self.btn_view_m2.setEnabled(False)
            self.btn_view_cross_profile.setEnabled(True)

            with open(self.param_file_path, 'r', encoding='utf-8') as f:
                params = json.load(f)

            self.spin_ma_size.blockSignals(True)
            self.spin_ma_size.setValue(params.get("ma_kernel_size", 31))
            self.spin_ma_size.blockSignals(False)

            self.spin_cross_size.setValue(params.get("cross_size", 40))
            c_min = params.get("colorbar_min", np.min(self.result_matrix))
            c_max = params.get("colorbar_max", np.max(self.result_matrix))

            self.click_points = [tuple(p) for p in params.get("heatmap_click_points", [])]
            if len(self.click_points) > 0:
                self.update_measure_display()
                self.redraw_measure_crosses()

            self.contour_click_points = [tuple(p) for p in params.get("contour_click_points", [])]

            max_val = np.max(self.result_matrix)
            peak_idx = np.unravel_index(np.argmax(self.result_matrix, axis=None), self.result_matrix.shape)

            self.lbl_status.setText("狀態: 已成功載入參數檔並重繪")
            self.lbl_status.setStyleSheet("color: #2E7D32; font-weight: bold;")
            self.lbl_max1.setText(f"Max Peak: {max_val:.1f}")
            self.lbl_size.setText(f"矩陣大小: {self.result_matrix.shape[0]} × {self.result_matrix.shape[1]}")

            self.image_item.setImage(self.result_matrix.T)
            self.hist.setLevels(self.result_matrix.min(), self.result_matrix.max())

            if "dr_shape_type" in params:
                if params["dr_shape_type"] == "ellipse":
                    self.radio_dr_shape_ellipse.setChecked(True)
                else:
                    self.radio_dr_shape_circle.setChecked(True)
            if "dr_enable_spot" in params:
                self.chk_dr_enable_spot.setChecked(bool(params["dr_enable_spot"]))
            if "dr_use_threshold" in params:
                self.chk_dr_use_threshold.setChecked(params["dr_use_threshold"])
            if "dr_threshold_percent" in params:
                self.spin_dr_thresh_percent.setValue(params["dr_threshold_percent"])
            if "m1_use_threshold" in params:
                self.chk_m1_use_threshold.setChecked(bool(params["m1_use_threshold"]))
            if "m1_threshold_percent" in params:
                self.spin_m1_thresh_percent.setValue(params["m1_threshold_percent"])
            if "p2_use_threshold" in params:
                self.chk_p2_use_threshold.setChecked(bool(params["p2_use_threshold"]))
            if "p2_threshold_percent" in params:
                self.spin_p2_thresh_percent.setValue(params["p2_threshold_percent"])
            if "dr_circle_diameter_px" in params:
                self.spin_dr_circle_diameter.setValue(params["dr_circle_diameter_px"])
            if "dr_ellipse_wx_px" in params:
                self.spin_dr_ellipse_wx.setValue(params["dr_ellipse_wx_px"])
            if "dr_ellipse_wy_px" in params:
                self.spin_dr_ellipse_wy.setValue(params["dr_ellipse_wy_px"])

            if "dr_center_x_px" in params and "dr_center_y_px" in params:
                self.dataray_center = (params["dr_center_x_px"], params["dr_center_y_px"])
            else:
                self.recalculate_dataray_spot()

            # 還原第一點與第二點的定位模式與門檻比例
            m1_mode = params.get("m1_point_mode", "centroid")
            p2_mode = params.get("p2_point_mode", "auto_min")
            
            for r in (self.radio_m1_peak_geom, self.radio_m1_centroid, self.radio_m1_thresh_geom, self.radio_m1_manual,
                      self.radio_p2_auto_min, self.radio_p2_m2_thresh_geom, self.radio_p2_m2_centroid, self.radio_p2_manual):
                r.blockSignals(True)
                
            if m1_mode == "centroid":
                self.radio_m1_centroid.setChecked(True)
            elif m1_mode == "thresh_geom":
                self.radio_m1_thresh_geom.setChecked(True)
            elif m1_mode == "peak_geom":
                self.radio_m1_peak_geom.setChecked(True)
            elif m1_mode == "manual":
                self.radio_m1_manual.setChecked(True)
                
            if p2_mode == "manual":
                self.radio_p2_manual.setChecked(True)
            elif p2_mode == "m2_thresh_geom":
                self.radio_p2_m2_thresh_geom.setChecked(True)
            elif p2_mode == "m2_centroid":
                self.radio_p2_m2_centroid.setChecked(True)
            else:
                self.radio_p2_auto_min.setChecked(True)
                
            for r in (self.radio_m1_peak_geom, self.radio_m1_centroid, self.radio_m1_thresh_geom, self.radio_m1_manual,
                      self.radio_p2_auto_min, self.radio_p2_m2_thresh_geom, self.radio_p2_m2_centroid, self.radio_p2_manual):
                r.blockSignals(False)

            if "m1_center_x_px" in params and "m1_center_y_px" in params and params["m1_center_x_px"] is not None:
                self.m1_center_point = (params["m1_center_x_px"], params["m1_center_y_px"])
                self.update_all_m1_markers()
            elif self.matrix1 is not None and not self.radio_m1_manual.isChecked():
                self.apply_m1_point_from_mode()

            if self._is_p2_auto_mode():
                self.sync_dual_points_after_m1_change()
            elif len(self.click_points) > 0:
                self._ensure_measure_mode_for_dual_points()
                self.update_measure_display()
                self.redraw_measure_crosses()

            self.render_contour_map(self.result_matrix)
            self.contour_hist.setLevels(c_min, c_max)
            self.render_sub_plots_fast(self.result_matrix, peak_idx[0])
            self.btn_export.setEnabled(True)
        except Exception as e:
            self.lbl_status.setText("狀態: 參數載入失敗")
            self.lbl_status.setStyleSheet("color: #C62828; font-weight: bold;")
            QMessageBox.critical(self, "錯誤", f"讀取參數或解析時發生錯誤：\n{str(e)}")

    def plot_single_file(self):
        if not self.file1_path:
            QMessageBox.warning(self, "警告", "請確保檔案已成功匯入！")
            return
        try:
            self.lbl_status.setText("狀態: 正在讀取與分析數據...")
            self.lbl_status.setStyleSheet("color: #F57C00; font-weight: bold;")
            self.repaint() # 取代 QApplication.processEvents()
            
            self.matrix1 = load_numeric_matrix(self.file1_path)
            self.matrix2 = None
            self.result_matrix = self.matrix1

            self.btn_view_m1.setEnabled(True)
            self.btn_view_m2.setEnabled(False)
            self.btn_view_cross_profile.setEnabled(True)

            max_val = np.max(self.result_matrix)
            peak_idx = np.unravel_index(np.argmax(self.result_matrix, axis=None), self.result_matrix.shape)
            
            self.lbl_status.setText("狀態: 畫圖成功")
            self.lbl_status.setStyleSheet("color: #2E7D32; font-weight: bold;")
            self.lbl_max1.setText(f"Max Peak: {max_val:.1f} (Row:{peak_idx[0]}, Col:{peak_idx[1]})")
            self.lbl_size.setText(f"矩陣大小: {self.result_matrix.shape[0]} × {self.result_matrix.shape[1]}")
            
            self.image_item.setImage(self.result_matrix.T)
            self.hist.setLevels(self.result_matrix.min(), self.result_matrix.max())
            
            self.recalculate_dataray_spot()
            self.apply_m1_point_from_mode()
            self.sync_dual_points_after_m1_change()
            self.render_contour_map(self.result_matrix)
            self.render_sub_plots_fast(self.result_matrix, peak_idx[0])
            self.btn_export.setEnabled(True)
        except Exception as e:
            self.lbl_status.setText("狀態: 讀取失敗")
            self.lbl_status.setStyleSheet("color: #C62828; font-weight: bold;")
            QMessageBox.critical(self, "錯誤", f"處理檔案時發生錯誤：\n{str(e)}")

    def plot_pure_subtraction(self):
        if not self.file1_path or not self.file2_path:
            QMessageBox.warning(self, "警告", "請確保兩個檔案都已成功匯入！")
            return
        try:
            self.lbl_status.setText("狀態: 正在進行純相減運算...")
            self.lbl_status.setStyleSheet("color: #F57C00; font-weight: bold;")
            self.repaint()
            
            m1 = load_numeric_matrix(self.file1_path)
            m2 = load_numeric_matrix(self.file2_path)
            
            if m1.shape != m2.shape:
                raise ValueError(f"兩個檔案的數據矩陣大小不一致！\n({m1.shape} vs {m2.shape})")
                
            self.matrix1 = m1
            self.matrix2 = m2

            if self.chk_normalize_peaks.isChecked():
                max1 = np.max(m1)
                max2 = np.max(m2)
                if max2 != 0:
                    scale = max1 / max2
                    m2_processed = m2 * scale
                else:
                    m2_processed = m2
            else:
                m2_processed = m2

            self.result_matrix = m1 - m2_processed

            self.btn_view_m1.setEnabled(True)
            self.btn_view_m2.setEnabled(True)
            self.btn_view_cross_profile.setEnabled(True)

            peak_idx = np.unravel_index(np.argmax(self.result_matrix, axis=None), self.result_matrix.shape)
            
            self.lbl_status.setText("狀態: 純相減畫圖成功" + (" (已強度對齊)" if self.chk_normalize_peaks.isChecked() else ""))
            self.lbl_status.setStyleSheet("color: #2E7D32; font-weight: bold;")
            self.lbl_max1.setText(f"Max 1: {np.max(m1):.1f}")
            self.lbl_max2.setText(f"Max 2: {np.max(m2):.1f}")
            self.lbl_size.setText(f"矩陣大小: {self.result_matrix.shape[0]} × {self.result_matrix.shape[1]}")
            
            self.image_item.setImage(self.result_matrix.T)
            self.hist.setLevels(self.result_matrix.min(), self.result_matrix.max())
            
            self.recalculate_dataray_spot()
            self.apply_m1_point_from_mode()
            self.sync_dual_points_after_m1_change()
            self.render_contour_map(self.result_matrix)
            self.render_sub_plots_fast(self.result_matrix, peak_idx[0])
            self.btn_export.setEnabled(True)
        except Exception as e:
            self.lbl_status.setText("狀態: 計算失敗")
            self.lbl_status.setStyleSheet("color: #C62828; font-weight: bold;")
            QMessageBox.critical(self, "錯誤", f"處理檔案時發生錯誤：\n{str(e)}")

    def plot_pure_division(self):
        if not self.file1_path or not self.file2_path:
            QMessageBox.warning(self, "警告", "請確保兩個檔案都已成功匯入！")
            return
        try:
            self.lbl_status.setText("狀態: 正在進行純相除運算...")
            self.lbl_status.setStyleSheet("color: #F57C00; font-weight: bold;")
            self.repaint()
            
            m1 = load_numeric_matrix(self.file1_path)
            m2 = load_numeric_matrix(self.file2_path)
            
            if m1.shape != m2.shape:
                raise ValueError(f"兩個檔案的數據矩陣大小不一致！\n({m1.shape} vs {m2.shape})")
                
            self.matrix1 = m1
            self.matrix2 = m2

            if self.chk_normalize_peaks.isChecked():
                max1 = np.max(m1)
                max2 = np.max(m2)
                if max2 != 0:
                    scale = max1 / max2
                    m2_processed = m2 * scale
                else:
                    m2_processed = m2
            else:
                m2_processed = m2

            safe_m2 = np.where(m2_processed == 0, 1e-9, m2_processed)
            self.result_matrix = m1 / safe_m2

            self.btn_view_m1.setEnabled(True)
            self.btn_view_m2.setEnabled(True)
            self.btn_view_cross_profile.setEnabled(True)

            peak_idx = np.unravel_index(np.argmax(self.result_matrix, axis=None), self.result_matrix.shape)
            
            self.lbl_status.setText("狀態: 純相除畫圖成功" + (" (已強度對齊)" if self.chk_normalize_peaks.isChecked() else ""))
            self.lbl_status.setStyleSheet("color: #2E7D32; font-weight: bold;")
            self.lbl_max1.setText(f"Max 1: {np.max(m1):.1f}")
            self.lbl_max2.setText(f"Max 2: {np.max(m2):.1f}")
            self.lbl_size.setText(f"矩陣大小: {self.result_matrix.shape[0]} × {self.result_matrix.shape[1]}")
            
            self.image_item.setImage(self.result_matrix.T)
            self.hist.setLevels(self.result_matrix.min(), self.result_matrix.max())
            
            self.recalculate_dataray_spot()
            self.apply_m1_point_from_mode()
            self.sync_dual_points_after_m1_change()
            self.render_contour_map(self.result_matrix)
            self.render_sub_plots_fast(self.result_matrix, peak_idx[0])
            self.btn_export.setEnabled(True)
        except Exception as e:
            self.lbl_status.setText("狀態: 計算失敗")
            self.lbl_status.setStyleSheet("color: #C62828; font-weight: bold;")
            QMessageBox.critical(self, "錯誤", f"處理檔案時發生錯誤：\n{str(e)}")

    def calculate_and_plot(self):
        if not self.file1_path or not self.file2_path:
            QMessageBox.warning(self, "警告", "請確保兩個檔案都已成功匯入！")
            return
        try:
            self.lbl_status.setText("狀態: 正在計算新方程式...")
            self.lbl_status.setStyleSheet("color: #F57C00; font-weight: bold;")
            self.repaint()
            
            self.matrix1 = load_numeric_matrix(self.file1_path)
            self.matrix2 = load_numeric_matrix(self.file2_path)
            
            if self.matrix1.shape != self.matrix2.shape:
                raise ValueError(f"兩個檔案的數據矩陣大小不一致！\n({self.matrix1.shape} vs {self.matrix2.shape})")
            
            max1_idx = np.unravel_index(np.argmax(self.matrix1, axis=None), self.matrix1.shape)
            max1_val = self.matrix1[max1_idx]
            match2_val = self.matrix2[max1_idx] if self.matrix2[max1_idx] != 0 else 1e-9
                
            scale_ratio = max1_val / match2_val
            self.result_matrix = self.matrix1 - (self.matrix2 * scale_ratio)

            self.btn_view_m1.setEnabled(True)
            self.btn_view_m2.setEnabled(True)
            self.btn_view_cross_profile.setEnabled(True)

            peak_idx = np.unravel_index(np.argmax(self.result_matrix, axis=None), self.result_matrix.shape)
            
            self.lbl_status.setText("狀態: 新方程式計算成功")
            self.lbl_status.setStyleSheet("color: #2E7D32; font-weight: bold;")
            self.lbl_max1.setText(f"M1最大值: {max1_val:.1f}")
            self.lbl_max2.setText(f"M2同位置: {match2_val:.1f}")
            self.lbl_ratio.setText(f"計算得出比例: {scale_ratio:.4f}")
            self.lbl_size.setText(f"矩陣大小: {self.result_matrix.shape[0]} × {self.result_matrix.shape[1]}")
            
            self.image_item.setImage(self.result_matrix.T)
            self.hist.setLevels(self.result_matrix.min(), self.result_matrix.max())
            
            self.recalculate_dataray_spot()
            self.apply_m1_point_from_mode()
            self.sync_dual_points_after_m1_change()
            self.render_contour_map(self.result_matrix)
            self.render_sub_plots_fast(self.result_matrix, peak_idx[0])
            self.btn_export.setEnabled(True)
        except Exception as e:
            self.lbl_status.setText("狀態: 計算失敗")
            self.lbl_status.setStyleSheet("color: #C62828; font-weight: bold;")
            QMessageBox.critical(self, "錯誤", f"處理檔案時發生錯誤：\n{str(e)}")

    def _create_hline(self):
        line = QFrame()
        line.setFrameShape(QFrame.HLine)
        line.setFrameShadow(QFrame.Sunken)
        line.setStyleSheet("color: #e0e0e0; margin-top: 2px; margin-bottom: 2px;")
        return line

    def on_ma_kernel_spin_changed(self):
        if self.result_matrix is not None:
            self.ma_timer.start(300)

    def apply_ma_kernel_change(self):
        if self.result_matrix is not None:
            self.render_contour_map(self.result_matrix)

    def on_contour_colorbar_levels_changed(self):
        self.cbar_timer.start(150)

    def update_isocurves_and_waveform(self):
        if self.smoothed_matrix is None:
            return
            
        min_v, max_v = self.contour_hist.getLevels()
        if min_v < max_v:
            levels = np.linspace(min_v, max_v, 10)
            if len(self.contour_curves) == len(levels):
                for idx, level in enumerate(levels):
                    self.contour_curves[idx].setLevel(level)
            else:
                for c_item in self.contour_curves:
                    self.plot_contour.removeItem(c_item)
                self.contour_curves.clear()
                
                for level in levels:
                    try:
                        isolines = pg.IsocurveItem(data=self.smoothed_matrix.T, level=level, pen=pg.mkPen('w', width=0.8))
                        self.plot_contour.addItem(isolines)
                        self.contour_curves.append(isolines)
                    except Exception:
                        pass
        self.redraw_contour_overlay()

    def clear_contour_lines(self):
        self.contour_click_points.clear()
        self.clear_contour_measure_items()
        self.plot_contour_profile.clear()

    def on_colorbar_levels_changed(self):
        if self.result_matrix is not None:
            min_lvl, max_lvl = self.hist.getLevels()
            self.plot_hist.setXRange(min_lvl, max_lvl, padding=0)

            if self.cross_profile_win is not None and self.cross_profile_win.isVisible():
                cx, cy = self.heatmap_cross_point if self.heatmap_cross_point is not None else (self.result_matrix.shape[1]//2, self.result_matrix.shape[0]//2)
                self.cross_profile_win.update_profiles(self.result_matrix, cx, cy, y_range=(min_lvl, max_lvl))